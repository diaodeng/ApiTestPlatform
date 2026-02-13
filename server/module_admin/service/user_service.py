import io
from typing import Union

from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from module_admin.dao.user_dao import UserDao
from module_admin.entity.vo.common_vo import CrudResponseModel, DataScopeExpr
from module_admin.entity.vo.user_vo import (
    AddUserModel,
    CrudUserRoleModel,
    CurrentUserModel,
    DeleteUserModel,
    EditUserModel,
    ResetUserModel,
    SelectedRoleModel,
    UserDetailModel,
    UserInfoModel,
    UserModel,
    UserPageQueryModel,
    UserPostModel,
    UserProfileModel,
    UserRoleModel,
    UserRoleQueryModel,
    UserRoleResponseModel,
)
from module_admin.service.post_service import PostPageQueryModel, PostService
from module_admin.service.role_service import RoleService
from utils.common_util import CamelCaseUtil, export_list2excel, get_excel_template
from utils.page_util import PageResponseModel
from utils.pwd_util import PwdUtil


class UserService:
    """
    用户管理模块服务层
    """

    @classmethod
    def get_user_list_services(cls,
                               query_db: Session,
                               query_object: UserPageQueryModel,
                               data_scope_sql: DataScopeExpr,
                               is_page: bool = False):
        """
        获取用户列表信息service
        :param query_db: orm对象
        :param query_object: 查询参数对象
        :param data_scope_sql: 数据权限对应的查询sql语句
        :param is_page: 是否开启分页
        :return: 用户列表信息对象
        """
        query_result = UserDao.get_user_list(query_db, query_object, data_scope_sql, is_page)
        if is_page:
            user_list_result = PageResponseModel(
                **{
                    **query_result.model_dump(by_alias=True),
                    'rows': [{**row[0], 'dept': row[1]} for row in query_result.rows]
                }
            )
        else:
            user_list_result = []
            if query_result:
                user_list_result = [{**row[0], 'dept': row[1]} for row in query_result]

        return user_list_result

    @classmethod
    def add_user_services(cls, query_db: Session, page_object: AddUserModel):
        """
        新增用户信息service
        :param query_db: orm对象
        :param page_object: 新增用户对象
        :return: 新增用户校验结果
        """
        add_user = UserModel(**page_object.model_dump(by_alias=True))
        user = UserDao.get_user_by_info(query_db, UserModel(userName=page_object.user_name))
        if user:
            result = {'is_success': False, 'message': '用户名已存在'}
        else:
            try:
                add_result = UserDao.add_user_dao(query_db, add_user)
                user_id = add_result.user_id
                if page_object.role_ids:
                    for role in page_object.role_ids:
                        UserDao.add_user_role_dao(query_db, UserRoleModel(userId=user_id, roleId=role))
                if page_object.post_ids:
                    for post in page_object.post_ids:
                        UserDao.add_user_post_dao(query_db, UserPostModel(userId=user_id, postId=post))
                query_db.commit()
                result = {'is_success': True, 'message': '新增成功'}
            except Exception as e:
                query_db.rollback()
                raise e

        return CrudResponseModel(**result)

    @classmethod
    def edit_user_services(cls, query_db: Session, page_object: EditUserModel):
        """
        编辑用户信息service
        :param query_db: orm对象
        :param page_object: 编辑用户对象
        :return: 编辑用户校验结果
        """
        edit_user = page_object.model_dump(exclude_unset=True, exclude={'admin'})
        if page_object.type != 'status' and page_object.type != 'avatar' and page_object.type != 'pwd':
            del edit_user['role_ids']
            del edit_user['post_ids']
            del edit_user['role']
        if page_object.type == 'status' or page_object.type == 'avatar' or page_object.type == 'pwd':
            del edit_user['type']
        user_info = cls.user_detail_services(query_db, edit_user.get('user_id'))
        if user_info:
            if (page_object.type != 'status'
                    and page_object.type != 'avatar'
                    and page_object.type == 'pwd'
                    and user_info.data.user_name != page_object.user_name):
                user = UserDao.get_user_by_info(query_db, UserModel(userName=page_object.user_name))
                if user:
                    result = {'is_success': False, 'message': '用户名已存在'}
                    return CrudResponseModel(**result)
            try:
                UserDao.edit_user_dao(query_db, edit_user)
                if page_object.type != 'status' and page_object.type != 'avatar':
                    UserDao.delete_user_role_dao(query_db, UserRoleModel(userId=page_object.user_id))
                    UserDao.delete_user_post_dao(query_db, UserPostModel(userId=page_object.user_id))
                    if page_object.role_ids:
                        for role in page_object.role_ids:
                            UserDao.add_user_role_dao(query_db, UserRoleModel(userId=page_object.user_id, roleId=role))
                    if page_object.post_ids:
                        for post in page_object.post_ids:
                            UserDao.add_user_post_dao(query_db, UserPostModel(userId=page_object.user_id, postId=post))
                query_db.commit()
                result = {'is_success': True, 'message': '更新成功'}
            except Exception as e:
                query_db.rollback()
                raise e
        else:
            result = {'is_success': False, 'message': '用户不存在'}

        return CrudResponseModel(**result)

    @classmethod
    def delete_user_services(cls, query_db: Session, page_object: DeleteUserModel):
        """
        删除用户信息service
        :param query_db: orm对象
        :param page_object: 删除用户对象
        :return: 删除用户校验结果
        """
        if page_object.user_ids.split(','):
            user_id_list = page_object.user_ids.split(',')
            try:
                for user_id in user_id_list:
                    user_id_dict = {'userId': user_id, 
                                    'updateBy': page_object.update_by, 
                                    'updateTime': page_object.update_time}
                    UserDao.delete_user_role_dao(query_db, UserRoleModel(**user_id_dict))
                    UserDao.delete_user_post_dao(query_db, UserPostModel(**user_id_dict))
                    UserDao.delete_user_dao(query_db, UserModel(**user_id_dict))
                query_db.commit()
                result = {'is_success': True, 'message': '删除成功'}
            except Exception as e:
                query_db.rollback()
                raise e
        else:
            result = {'is_success': False, 'message': '传入用户id为空'}
        return CrudResponseModel(**result)

    @classmethod
    def user_detail_services(cls, query_db: Session, user_id: Union[int, str]):
        """
        获取用户详细信息service
        :param query_db: orm对象
        :param user_id: 用户id
        :return: 用户id对应的信息
        """
        posts = PostService.get_post_list_services(query_db, PostPageQueryModel(**{}), is_page=False)
        roles = RoleService.get_role_select_option_services(query_db)
        if user_id != '':
            query_user = UserDao.get_user_detail_by_id(query_db, user_id=user_id)
            post_ids = ','.join([str(row.post_id) for row in query_user.get('user_post_info')])
            post_ids_list = [row.post_id for row in query_user.get('user_post_info')]
            role_ids = ','.join([str(row.role_id) for row in query_user.get('user_role_info')])
            role_ids_list = [row.role_id for row in query_user.get('user_role_info')]

            return UserDetailModel(
                data=UserInfoModel(
                    **CamelCaseUtil.transform_result(query_user.get('user_basic_info')),
                    postIds=post_ids,
                    roleIds=role_ids,
                    dept=CamelCaseUtil.transform_result(query_user.get('user_dept_info')),
                    role=CamelCaseUtil.transform_result(query_user.get('user_role_info'))
                ),
                postIds=post_ids_list,
                posts=posts,
                roleIds=role_ids_list,
                roles=roles
            )

        return UserDetailModel(
            posts=posts,
            roles=roles
        )

    @classmethod
    def user_profile_services(cls, query_db: Session, user_id: int):
        """
        获取用户详细信息service
        :param query_db: orm对象
        :param user_id: 用户id
        :return: 用户id对应的信息
        """
        query_user = UserDao.get_user_detail_by_id(query_db, user_id=user_id)
        post_ids = ','.join([str(row.post_id) for row in query_user.get('user_post_info')])
        post_group = ','.join([row.post_name for row in query_user.get('user_post_info')])
        role_ids = ','.join([str(row.role_id) for row in query_user.get('user_role_info')])
        role_group = ','.join([row.role_name for row in query_user.get('user_role_info')])

        return UserProfileModel(
            data=UserInfoModel(
                **CamelCaseUtil.transform_result(query_user.get('user_basic_info')),
                postIds=post_ids,
                roleIds=role_ids,
                dept=CamelCaseUtil.transform_result(query_user.get('user_dept_info')),
                role=CamelCaseUtil.transform_result(query_user.get('user_role_info'))
            ),
            postGroup=post_group,
            roleGroup=role_group
        )

    @classmethod
    def reset_user_services(cls, query_db: Session, page_object: ResetUserModel):
        """
        重置用户密码service
        :param query_db: orm对象
        :param page_object: 重置用户对象
        :return: 重置用户校验结果
        """
        reset_user = page_object.model_dump(exclude_unset=True, exclude={'admin'})
        if page_object.old_password:
            user = UserDao.get_user_detail_by_id(query_db, user_id=page_object.user_id).get('user_basic_info')
            if not PwdUtil.verify_password(page_object.old_password, user.password):
                result = {'is_success': False, 'message': '旧密码不正确'}
                return CrudResponseModel(**result)
            else:
                del reset_user['old_password']
        if page_object.sms_code and page_object.session_id:
            del reset_user['sms_code']
            del reset_user['session_id']
        try:
            UserDao.edit_user_dao(query_db, reset_user)
            query_db.commit()
            result = {"is_success": True, "message": "重置成功"}
        except Exception as e:
            query_db.rollback()
            raise e

        return CrudResponseModel(**result)

    @classmethod
    async def batch_import_user_services(
        cls, query_db: Session, file: UploadFile, update_support: bool, current_user: CurrentUserModel
    ):
        """
        批量导入用户service
        """

        header_dict = {
            "部门编号": "dept_id",
            "登录名称": "user_name",
            "用户名称": "nick_name",
            "用户邮箱": "email",
            "手机号码": "phonenumber",
            "用户性别": "sex",
            "帐号状态": "status",
        }

        contents = await file.read()
        await file.close()

        wb = load_workbook(io.BytesIO(contents))
        ws = wb.active  # 默认第一个sheet

        # 读取表头
        headers = [cell.value for cell in ws[1]]
        mapped_headers = [header_dict.get(h) for h in headers]

        add_error_result = []
        count = 0

        try:
            # 从第二行开始遍历
            for row in ws.iter_rows(min_row=2, values_only=True):
                count += 1

                row_dict = dict(zip(mapped_headers, row))

                if not row_dict.get("user_name"):
                    continue  # 跳过空行

                # 数据转换
                sex_map = {"男": "0", "女": "1", "未知": "2"}
                status_map = {"正常": "0", "停用": "1"}

                row_dict["sex"] = sex_map.get(row_dict.get("sex"))
                row_dict["status"] = status_map.get(row_dict.get("status"))

                user_info = UserDao.get_user_by_info(query_db, UserModel(userName=row_dict.get("user_name")))

                if user_info:
                    if update_support:
                        edit_user = UserModel(
                            userId=user_info.user_id,
                            deptId=row_dict.get("dept_id"),
                            userName=row_dict.get("user_name"),
                            nickName=row_dict.get("nick_name"),
                            email=row_dict.get("email"),
                            phonenumber=str(row_dict.get("phonenumber")),
                            sex=row_dict.get("sex"),
                            status=row_dict.get("status"),
                            updateBy=current_user.user.user_name,
                        ).model_dump(exclude_unset=True)

                        UserDao.edit_user_dao(query_db, edit_user)
                    else:
                        add_error_result.append(f"{count}.用户账号{row_dict.get('user_name')}已存在")
                else:
                    add_user = UserModel(
                        deptId=row_dict.get("dept_id"),
                        userName=row_dict.get("user_name"),
                        password=PwdUtil.get_password_hash("123456"),
                        nickName=row_dict.get("nick_name"),
                        email=row_dict.get("email"),
                        phonenumber=str(row_dict.get("phonenumber")),
                        sex=row_dict.get("sex"),
                        status=row_dict.get("status"),
                        createBy=current_user.user.user_name,
                        updateBy=current_user.user.user_name,
                    )
                    UserDao.add_user_dao(query_db, add_user)

            query_db.commit()

            result = {
                "is_success": True,
                "message": "\n".join(add_error_result)
            }

        except Exception as e:
            query_db.rollback()
            raise e

        return CrudResponseModel(**result)


    @staticmethod
    def get_user_import_template_services():
        """
        获取用户导入模板service
        :return: 用户导入模板excel的二进制数据
        """
        header_list = ["部门编号", "登录名称", "用户名称", "用户邮箱", "手机号码", "用户性别", "帐号状态"]
        selector_header_list = ["用户性别", "帐号状态"]
        option_list = [{"用户性别": ["男", "女", "未知"]}, {"帐号状态": ["正常", "停用"]}]
        binary_data = get_excel_template(header_list=header_list,
                                         selector_header_list=selector_header_list,
                                         option_list=option_list)

        return binary_data

    @staticmethod
    def export_user_list_services(user_list: list):
        """
        导出用户信息service
        :param user_list: 用户信息列表
        :return: 用户信息对应excel的二进制数据
        """
        # 创建一个映射字典，将英文键映射到中文键
        mapping_dict = {
            "userId": "用户编号",
            "userName": "用户名称",
            "nickName": "用户昵称",
            "deptName": "部门",
            "email": "邮箱地址",
            "phonenumber": "手机号码",
            "sex": "性别",
            "status": "状态",
            "createBy": "创建者",
            "createTime": "创建时间",
            "updateBy": "更新者",
            "updateTime": "更新时间",
            "remark": "备注",
        }

        data = user_list

        for item in data:
            if item.get('status') == '0':
                item['status'] = '正常'
            else:
                item['status'] = '停用'
            if item.get('sex') == '0':
                item['sex'] = '男'
            elif item.get('sex') == '1':
                item['sex'] = '女'
            else:
                item['sex'] = '未知'
        new_data = [{mapping_dict.get(key): value for key, value in item.items() if mapping_dict.get(key)} for item in data]
        binary_data = export_list2excel(new_data)

        return binary_data

    @classmethod
    def get_user_role_allocated_list_services(cls, query_db: Session, page_object: UserRoleQueryModel):
        """
        根据用户id获取已分配角色列表
        :param query_db: orm对象
        :param page_object: 用户关联角色对象
        :return: 已分配角色列表
        """
        query_user = UserDao.get_user_detail_by_id(query_db, page_object.user_id)
        post_ids = ','.join([str(row.post_id) for row in query_user.get('user_post_info')])
        role_ids = ','.join([str(row.role_id) for row in query_user.get('user_role_info')])
        user = UserInfoModel(
            **CamelCaseUtil.transform_result(query_user.get('user_basic_info')),
            postIds=post_ids,
            roleIds=role_ids,
            dept=CamelCaseUtil.transform_result(query_user.get('user_dept_info')),
            role=CamelCaseUtil.transform_result(query_user.get('user_role_info'))
        )
        query_role_list = [SelectedRoleModel(**row) for row in RoleService.get_role_select_option_services(query_db)]
        for model_a in query_role_list:
            for model_b in user.role:
                if model_a.role_id == model_b.role_id:
                    model_a.flag = True
        result = UserRoleResponseModel(
            roles=query_role_list,
            user=user
        )

        return result

    @classmethod
    def add_user_role_services(cls, query_db: Session, page_object: CrudUserRoleModel):
        """
        新增用户关联角色信息service
        :param query_db: orm对象
        :param page_object: 新增用户关联角色对象
        :return: 新增用户关联角色校验结果
        """
        if page_object.user_id and page_object.role_ids:
            role_id_list = page_object.role_ids.split(',')
            try:
                for role_id in role_id_list:
                    user_role = cls.detail_user_role_services(query_db, UserRoleModel(userId=page_object.user_id, roleId=role_id))
                    if user_role:
                        continue
                    else:
                        UserDao.add_user_role_dao(query_db, UserRoleModel(userId=page_object.user_id, roleId=role_id))
                query_db.commit()
                result = {'is_success': True, 'message': '分配成功'}
            except Exception as e:
                query_db.rollback()
                raise e
        elif page_object.user_id and not page_object.role_ids:
            try:
                UserDao.delete_user_role_by_user_and_role_dao(query_db, UserRoleModel(userId=page_object.user_id))
                query_db.commit()
                result = {'is_success': True, 'message': '分配成功'}
            except Exception as e:
                query_db.rollback()
                raise e
        elif page_object.user_ids and page_object.role_id:
            user_id_list = page_object.user_ids.split(',')
            try:
                for user_id in user_id_list:
                    user_role = cls.detail_user_role_services(query_db, UserRoleModel(userId=user_id, roleId=page_object.role_id))
                    if user_role:
                        continue
                    else:
                        UserDao.add_user_role_dao(query_db, UserRoleModel(userId=user_id, roleId=page_object.role_id))
                query_db.commit()
                result = {'is_success': True, 'message': '新增成功'}
            except Exception as e:
                query_db.rollback()
                raise e
        else:
            result = {'is_success': False, 'message': '不满足新增条件'}

        return CrudResponseModel(**result)

    @classmethod
    def delete_user_role_services(cls, query_db: Session, page_object: CrudUserRoleModel):
        """
        删除用户关联角色信息service
        :param query_db: orm对象
        :param page_object: 删除用户关联角色对象
        :return: 删除用户关联角色校验结果
        """
        if (page_object.user_id and page_object.role_id) or (page_object.user_ids and page_object.role_id):
            if page_object.user_id and page_object.role_id:
                try:
                    UserDao.delete_user_role_by_user_and_role_dao(query_db, UserRoleModel(userId=page_object.user_id, roleId=page_object.role_id))
                    query_db.commit()
                    result = {'is_success': True, 'message': '删除成功'}
                except Exception as e:
                    query_db.rollback()
                    raise e
            elif page_object.user_ids and page_object.role_id:
                user_id_list = page_object.user_ids.split(',')
                try:
                    for user_id in user_id_list:
                        UserDao.delete_user_role_by_user_and_role_dao(query_db, UserRoleModel(userId=user_id, roleId=page_object.role_id))
                    query_db.commit()
                    result = {'is_success': True, 'message': '删除成功'}
                except Exception as e:
                    query_db.rollback()
                    raise e
            else:
                result = {'is_success': False, 'message': '不满足删除条件'}
        else:
            result = {'is_success': False, 'message': '传入用户角色关联信息为空'}

        return CrudResponseModel(**result)

    @classmethod
    def detail_user_role_services(cls, query_db: Session, page_object: UserRoleModel):
        """
        获取用户关联角色详细信息service
        :param query_db: orm对象
        :param page_object: 用户关联角色对象
        :return: 用户关联角色详细信息
        """
        user_role = UserDao.get_user_role_detail(query_db, page_object)

        return user_role
