"""门店配置导入环境归属测试。

覆盖导入门店配置时的环境语义：
- 未选择环境时拒绝导入；
- 选择环境后，导入行数据归属该环境；
- 覆盖导入只清空所选环境旧数据，其他环境保持不变；
- 增量导入按 环境+商户+机构+SAP 四元组匹配去重。
"""
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook

from modules.ticket.dao.ticket_log_pull_dao import TicketLogPullDao
from modules.ticket.service.log_pull.ticket_log_pull_service import TicketLogPullService


class DummyUser:
    """仅用于通过类型参数传递的简化用户对象。"""

    user = SimpleNamespace(user_name="tester", user_id=7)


def _build_store_workbook() -> bytes:
    """构造一份带表头与两行门店数据的导入 Excel。"""
    workbook = Workbook()
    sheet = workbook.active
    for index, header in enumerate(TicketLogPullService.STORE_IMPORT_HEADERS, 1):
        sheet.cell(row=1, column=index, value=header)
    sample = [
        ["G001", "V10001", "R001", "O10001", "上海门店", "SAP10001"],
        ["G001", "V10001", "R001", "O10002", "北京门店", "SAP10002"],
    ]
    for row_offset, row_data in enumerate(sample, 2):
        for index, value in enumerate(row_data, 1):
            sheet.cell(row=row_offset, column=index, value=value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_import_store_config_requires_environment():
    """未选择环境时导入必须直接拒绝，不允许写入无归属数据。"""
    result = TicketLogPullService.import_store_config_services(
        object(), _build_store_workbook(), "incremental", DummyUser(), environment=""
    )
    assert result.is_success is False
    assert "环境" in result.message


def test_import_store_config_writes_environment_to_rows(monkeypatch):
    """导入成功后每行数据应归属所选环境。"""
    created_entities = []
    deleted_environments = []

    monkeypatch.setattr(
        TicketLogPullDao,
        "delete_store_configs_by_environment",
        lambda db, environment: deleted_environments.append(environment) or 0,
    )
    monkeypatch.setattr(TicketLogPullDao, "list_all_store_configs", lambda db: [])

    original_entity_builder = TicketLogPullService._build_store_config_entity.__func__

    def capture_entity(cls, row, now, *, environment=""):
        store = original_entity_builder(cls, row, now, environment=environment)
        created_entities.append(store)
        return store

    monkeypatch.setattr(
        TicketLogPullService, "_build_store_config_entity", classmethod(capture_entity)
    )

    captured_entities = []
    real_add = TicketLogPullDao.save_store_config

    def fake_add(db, store):
        captured_entities.append(store)
        return store

    # 直接给 session 加一个 add 记录器，service 内部使用 query_db.add(store)
    session = SimpleNamespace(
        add=lambda store: captured_entities.append(store),
        commit=lambda: None,
        rollback=lambda: None,
    )

    result = TicketLogPullService.import_store_config_services(
        session, _build_store_workbook(), "incremental", DummyUser(), environment="uat"
    )

    assert result.is_success is True
    assert result.result["insertedCount"] == 2
    assert result.result["environment"] == "uat"
    assert created_entities and all(store.environment == "uat" for store in created_entities)
    assert captured_entities and all(store.environment == "uat" for store in captured_entities)
    assert deleted_environments == []  # 增量导入不清空数据
    assert real_add is not None  # 保持引用避免未使用告警


def test_import_store_config_overwrite_only_clears_selected_environment(monkeypatch):
    """覆盖导入只清空所选环境，不得清空其他环境配置。"""
    deleted_environments = []

    monkeypatch.setattr(
        TicketLogPullDao,
        "delete_store_configs_by_environment",
        lambda db, environment: deleted_environments.append(environment) or 3,
    )
    monkeypatch.setattr(
        TicketLogPullDao,
        "delete_all_store_configs",
        lambda db: (_ for _ in ()).throw(AssertionError("覆盖导入不应清空全表")),
    )
    monkeypatch.setattr(TicketLogPullDao, "list_all_store_configs", lambda db: [])

    session = SimpleNamespace(add=lambda store: None, commit=lambda: None, rollback=lambda: None)

    result = TicketLogPullService.import_store_config_services(
        session, _build_store_workbook(), "overwrite", DummyUser(), environment="prod"
    )

    assert result.is_success is True
    assert deleted_environments == ["prod"]


def test_store_config_match_key_includes_environment():
    """去重键必须包含环境，保证同一门店可在不同环境各存一条。"""
    store = SimpleNamespace(
        environment="uat",
        vender_no="V10001",
        org_no="O10001",
        sap_org_no="SAP10001",
    )
    match_key = TicketLogPullService._build_store_config_match_key(store)
    assert match_key == ("uat", "V10001", "O10001", "SAP10001")

    store_other_env = SimpleNamespace(
        environment="prod",
        vender_no="V10001",
        org_no="O10001",
        sap_org_no="SAP10001",
    )
    assert TicketLogPullService._build_store_config_match_key(
        store_other_env
    ) != TicketLogPullService._build_store_config_match_key(store)


def test_store_config_entity_defaults_environment_empty():
    """构建实体时未传环境应为空串，与存量数据语义一致。"""
    store = TicketLogPullService._build_store_config_entity(
        {
            "vender_no": "V10001",
            "org_no": "O10001",
        },
        datetime.now(),
    )
    assert store.environment == ""
