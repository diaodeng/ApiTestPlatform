from __future__ import annotations

import base64
import gzip
import io
import json
import os
import re
import shutil
import tempfile
import threading
import traceback
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, time, timedelta
from ftplib import FTP, error_perm
from ftplib import all_errors as ftp_errors
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import urlparse

import requests
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from config.database import SessionLocal
from module_admin.entity.vo.user_vo import CurrentUserModel
from module_hrm.entity.vo.common_vo import CrudResponseModel
from modules.ticket.dao.ticket_dao import TicketDao
from modules.ticket.dao.ticket_log_pull_dao import TicketLogPullDao
from modules.ticket.entity.do.ticket_do import TicketEvent
from modules.ticket.entity.do.ticket_log_pull_do import (
    TicketLogPullProjectVendorMap,
    TicketLogPullRecord,
    TicketLogPullStoreConfig,
)
from modules.ticket.entity.vo.ticket_log_pull_vo import (
    TicketLogPullContentModel,
    TicketLogPullContentQueryModel,
    TicketLogPullCreateModel,
    TicketLogPullListItemModel,
    TicketLogPullProjectVendorMapModel,
    TicketLogPullProjectVendorMapQueryModel,
    TicketLogPullProjectVendorMapUpsertModel,
    TicketLogPullQueryModel,
    TicketLogPullStorageConfigModel,
    TicketLogPullStoreConfigQueryModel,
    TicketLogPullStoreOptionModel,
    TicketLogPullSummaryModel,
    TicketLogPullVendorOptionModel,
    TicketLogPullVendorStoreOptionsModel,
)
from modules.ticket.enums.ticket_enums import TicketEventType, TicketLogDataType, TicketLogPullStatus
from modules.ticket.service.notification.ticket_notify_service import TicketNotifyService
from utils.common_util import CamelCaseUtil
from utils.log_util import logger


class TicketLogContentTooLargeError(Exception):
    """
    日志内容超过入库上限时抛出的异常。
    """


class TicketLogPullService:
    """
    工单日志拉取服务层，负责配置管理、记录管理和后台拉取处理。
    """

    DEFAULT_INSERT_URL = ""
    DEFAULT_PAGE_URL = ""
    DEFAULT_TEMP_DIR = Path(__file__).resolve().parents[4] / "logs" / "ticket_log_pull" / "tmp"
    DEFAULT_LOCAL_DIR = Path(__file__).resolve().parents[4] / "logs" / "ticket_log_pull" / "archive"
    ACTIVE_STATUSES = {
        TicketLogPullStatus.CREATED.value,
        TicketLogPullStatus.SUBMITTING.value,
        TicketLogPullStatus.POLLING.value,
        TicketLogPullStatus.DOWNLOADING.value,
        TicketLogPullStatus.PROCESSING.value,
    }
    TIMESTAMP_PATTERN = re.compile(r"^(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2},\d{3})")
    _executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="ticket-log-pull")
    _executor_lock = threading.Lock()
    _active_record_ids: set[int] = set()
    VERSION_PATTERN = re.compile(r"(?:版本号|版本|version|app[_\s-]*version)[:：\s-]*([A-Za-z0-9._/-]+)", re.IGNORECASE)
    STORE_IMPORT_HEADERS = [
        "集团编号",
        "商户编号",
        "区域编号",
        "机构编号",
        "机构名称",
        "SAP机构编号",
        "会员渠道编号",
        "上级机构编号",
        "权限树节点ID",
        "机构类型",
        "所属公司代码",
        "城市编号",
        "业态编号",
        "状态",
        "创建时间",
        "修改时间",
        "开业日期",
        "默认语言",
    ]
    STORE_IMPORT_SAMPLE = {
        "集团编号": "G001",
        "商户编号": "V10001",
        "区域编号": "R001",
        "机构编号": "O10001",
        "机构名称": "示例门店",
        "SAP机构编号": "SAP10001",
        "会员渠道编号": "01",
        "上级机构编号": "P1000",
        "权限树节点ID": 0,
        "机构类型": 4,
        "所属公司代码": "C001",
        "城市编号": "010",
        "业态编号": "1",
        "状态": 1,
        "创建时间": "2026-01-01 10:00:00",
        "修改时间": "2026-01-01 10:00:00",
        "开业日期": "2026-01-01",
        "默认语言": "zh_HK",
    }

    @staticmethod
    def _user_id(current_user: CurrentUserModel) -> int | None:
        """
        获取当前登录用户ID。
        :param current_user: 当前登录用户
        :return: 用户ID
        """
        return current_user.user.user_id if current_user and current_user.user else None

    @staticmethod
    def _user_name(current_user: CurrentUserModel) -> str:
        """
        获取当前登录用户名。
        :param current_user: 当前登录用户
        :return: 用户名
        """
        if not current_user or not current_user.user:
            return "system"
        return current_user.user.user_name or current_user.user.nick_name or "system"

    @staticmethod
    def _json_dumps(value: Any) -> str:
        """
        将对象序列化为 JSON 字符串。
        :param value: 任意对象
        :return: JSON 字符串
        """
        return json.dumps(value, ensure_ascii=False)

    @staticmethod
    def _json_loads(value: Any, default: Any):
        """
        安全解析 JSON 字符串。
        :param value: 原始值
        :param default: 解析失败时默认值
        :return: 解析后的结果
        """
        if value in (None, ""):
            return default
        if isinstance(value, (dict, list)):
            return value
        try:
            return json.loads(str(value))
        except Exception:
            return default

    @staticmethod
    def _cell_text(value: Any) -> str:
        """
        将单元格值转为文本。
        :param value: 原始单元格值
        :return: 文本值
        """
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.isoformat(sep=" ")
        if isinstance(value, date):
            return value.isoformat()
        return str(value).strip()

    @staticmethod
    def _parse_import_datetime(value: Any) -> datetime | None:
        """
        解析导入文件中的日期时间值。
        :param value: 原始值
        :return: 日期时间
        """
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, time.min)
        text = str(value).strip()
        if not text:
            return None
        for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(text, pattern)
                if pattern == "%Y-%m-%d":
                    return datetime.combine(parsed.date(), time.min)
                return parsed
            except ValueError:
                continue
        return None

    @staticmethod
    def _parse_import_date(value: Any) -> date | None:
        """
        解析导入文件中的日期值。
        :param value: 原始值
        :return: 日期
        """
        if value in (None, ""):
            return None
        if isinstance(value, date) and not isinstance(value, datetime):
            return value
        if isinstance(value, datetime):
            return value.date()
        text = str(value).strip()
        if not text:
            return None
        for pattern in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _parse_import_int(value: Any, default: int = 0) -> int:
        """
        解析导入文件中的整数。
        :param value: 原始值
        :param default: 失败时默认值
        :return: 整数
        """
        if value in (None, ""):
            return default
        try:
            return int(float(str(value).strip()))
        except Exception:
            return default

    @staticmethod
    def _first_present_value(source: dict[str, Any], *keys: str) -> Any:
        """
        按字段优先级读取第一个非空值。
        :param source: 数据来源字典
        :param keys: 候选字段名列表
        :return: 第一个非空字段值，未命中时返回 None
        """
        for key in keys:
            value = source.get(key)
            if value not in (None, ""):
                return value
        return None

    @staticmethod
    def _parse_positive_int(value: Any, default: int) -> int:
        """
        将历史记录中的数值字段安全解析为正整数。
        :param value: 原始值
        :param default: 解析失败或非正数时的默认值
        :return: 正整数
        """
        try:
            parsed_value = int(float(str(value).strip()))
            return parsed_value if parsed_value > 0 else default
        except Exception:
            return default

    @classmethod
    def _log_chain_step(
        cls,
        query_db: Session,
        *,
        ticket_id: int | None,
        record_id: int | None,
        step: str,
        status: str,
        reason: str = "",
        detail: dict[str, Any] | None = None,
        operator_name: str = "system",
    ) -> None:
        """
        记录日志拉取链路步骤日志。
        :param query_db: 数据库会话
        :param ticket_id: 工单ID
        :param record_id: 日志拉取记录ID
        :param step: 当前步骤
        :param status: 步骤状态
        :param reason: 说明或跳过原因
        :param detail: 结构化详情
        :param operator_name: 操作人名称
        :return: 无
        """
        payload = {
            "record_id": record_id,
            "step": step,
            "status": status,
            "reason": reason,
            "detail": detail or {},
        }
        logger.info("日志拉取步骤: {}", payload)
        if ticket_id:
            cls._add_ticket_event(
                query_db,
                ticket_id=ticket_id,
                operator_id=None,
                operator_name=operator_name,
                content=f"{step}:{status}",
                event_data=payload,
            )

    @classmethod
    def _default_storage_config(cls) -> dict[str, Any]:
        """
        构建默认日志拉取存储配置。
        :return: 默认配置字典
        """
        return {
            "mode": "local",
            "localDirectory": "",
            "ftp": {
                "host": "",
                "port": 21,
                "username": "",
                "password": "",
                "baseDir": "",
                "passive": True,
                "timeoutSec": 15,
                "encoding": "utf-8",
            },
            "effectiveLocalDirectory": str(cls.DEFAULT_LOCAL_DIR.resolve()),
            "pollIntervalSec": 20,
            "pollTimeoutSec": 1800,
            "downloadTimeoutSec": 300,
            "maxContentChars": 500000,
        }

    @classmethod
    def _default_external_config(cls) -> dict[str, Any]:
        """
        构建默认日志拉取外部接口配置。
        :return: 默认外部接口配置
        """
        return {
            "insertUrl": cls.DEFAULT_INSERT_URL,
            "pageUrl": cls.DEFAULT_PAGE_URL,
            "headers": {
                "cookie": "",
                "origin": "https://erp.rta-os.com",
            },
            "vendors": [],
        }

    @classmethod
    def _default_parameter_examples(cls) -> list[dict[str, str]]:
        """
        构建日志拉取参数示例的默认配置。
        :return: 参数示例列表，元素包含 name/value
        """
        return [
            {"name": "modifyTime 日期示例", "value": "2026-06-23"},
            {"name": "path 路径示例", "value": "/path/to/log_or_database"},
        ]

    @classmethod
    def _resolve_local_dir(cls, raw_path: str | None = None) -> Path:
        """
        解析本地归档目录。
        :param raw_path: 用户配置目录
        :return: 绝对路径
        """
        raw_value = str(raw_path or "").strip()
        if not raw_value:
            return cls.DEFAULT_LOCAL_DIR.resolve()
        path = Path(raw_value).expanduser()
        if not path.is_absolute():
            path = Path(__file__).resolve().parents[4] / path
        return path.resolve()

    @classmethod
    def _normalize_storage_config(cls, raw_config: Any) -> dict[str, Any]:
        """
        归一化日志拉取存储配置。
        :param raw_config: 原始配置对象
        :return: 标准化后的配置字典
        """
        defaults = cls._default_storage_config()
        config = dict(raw_config or {}) if isinstance(raw_config, dict) else {}
        ftp_config = config.get("ftp") if isinstance(config.get("ftp"), dict) else {}
        mode = str(config.get("mode") or defaults["mode"]).strip().lower()
        if mode not in {"local", "ftp"}:
            mode = "local"
        normalized = {
            **defaults,
            **config,
            "mode": mode,
            "ftp": {**defaults["ftp"], **ftp_config},
        }
        normalized["pollIntervalSec"] = max(int(normalized.get("pollIntervalSec") or 20), 3)
        normalized["pollTimeoutSec"] = max(int(normalized.get("pollTimeoutSec") or 1800), 60)
        normalized["downloadTimeoutSec"] = max(int(normalized.get("downloadTimeoutSec") or 300), 30)
        normalized["maxContentChars"] = max(int(normalized.get("maxContentChars") or 500000), 10000)
        normalized["effectiveLocalDirectory"] = str(cls._resolve_local_dir(normalized.get("localDirectory")))
        return normalized

    @classmethod
    def _normalize_external_config(cls, raw_config: Any) -> dict[str, Any]:
        """
        归一化日志拉取外部接口配置。
        :param raw_config: 原始配置对象
        :return: 标准化后的外部接口配置
        """
        defaults = cls._default_external_config()
        config = dict(raw_config or {}) if isinstance(raw_config, dict) else {}
        headers = config.get("headers") if isinstance(config.get("headers"), dict) else {}
        normalized = {
            **defaults,
            **config,
            "headers": {**defaults["headers"], **headers},
        }
        normalized["insertUrl"] = str(normalized.get("insertUrl") or defaults["insertUrl"]).strip()
        normalized["pageUrl"] = str(normalized.get("pageUrl") or defaults["pageUrl"]).strip()
        normalized["headers"] = {
            key: str(value or "").strip()
            for key, value in normalized["headers"].items()
            if str(value or "").strip()
        }
        normalized["vendors"] = cls._normalize_vendor_store_options(config.get("vendors"))
        return normalized

    @classmethod
    def _normalize_vendor_store_options(cls, raw_vendors: Any) -> list[dict[str, Any]]:
        """
        归一化日志拉取商家/门店联动配置。
        :param raw_vendors: 原始商家配置列表
        :return: 规范化后的商家列表
        """
        if not isinstance(raw_vendors, list):
            return []

        normalized_vendors: list[dict[str, Any]] = []
        seen_vendor_ids: set[int] = set()
        for vendor in raw_vendors:
            if not isinstance(vendor, dict):
                continue
            try:
                vendor_id = int(vendor.get("vendorId"))
            except (TypeError, ValueError):
                continue
            if vendor_id <= 0 or vendor_id in seen_vendor_ids:
                continue
            seen_vendor_ids.add(vendor_id)
            stores = vendor.get("stores") if isinstance(vendor.get("stores"), list) else []
            normalized_stores: list[dict[str, Any]] = []
            seen_store_ids: set[str] = set()
            for store in stores:
                if not isinstance(store, dict):
                    continue
                store_id = cls._normalize_store_option_value(store.get("storeId"), store.get("storeCode"))
                if not store_id or store_id in seen_store_ids:
                    continue
                seen_store_ids.add(store_id)
                normalized_stores.append(
                    {
                        "storeId": store_id,
                        "storeCode": str(store.get("storeCode") or "").strip() or None,
                        "sapOrgNo": str(store.get("sapOrgNo") or "").strip() or None,
                        "storeName": str(store.get("storeName") or "").strip() or str(store_id),
                    }
                )
            normalized_vendors.append(
                {
                    "vendorId": vendor_id,
                    "vendorCode": str(vendor.get("vendorCode") or "").strip() or None,
                    "vendorName": str(vendor.get("vendorName") or "").strip() or str(vendor_id),
                    "stores": normalized_stores,
                }
            )
        return normalized_vendors

    @classmethod
    def _normalize_parameter_examples(cls, raw_examples: Any) -> list[dict[str, str]]:
        """
        归一化日志拉取参数示例配置。
        :param raw_examples: 参数配置原始值，期望为 name/value 字典列表
        :return: 过滤空值并去重后的示例列表
        """
        if not isinstance(raw_examples, list):
            return []
        normalized_examples: list[dict[str, str]] = []
        seen_values: set[str] = set()
        for item in raw_examples:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name") or item.get("label") or "").strip()
            value = str(item.get("value") or "").strip()
            if not value or value in seen_values:
                continue
            seen_values.add(value)
            normalized_examples.append({"name": name or value, "value": value})
        return normalized_examples

    @staticmethod
    def _extract_option_id(*values: Any) -> int | None:
        """
        从门店配置字段中提取联动选项ID。
        :param values: 候选值，优先使用可直接转成整数的值，其次提取文本中的数字
        :return: 可用于前端联动的整数ID
        """
        for value in values:
            if value in (None, ""):
                continue
            text = str(value).strip()
            if not text:
                continue
            try:
                return int(text)
            except (TypeError, ValueError):
                match = re.search(r"\d+", text)
                if match:
                    return int(match.group(0))
        return None

    @staticmethod
    def _normalize_store_option_value(*values: Any) -> str:
        """
        归一化门店下拉实际提交值，优先返回 org_no，其次回退到 SAP 编号。
        :param values: 候选门店标识
        :return: 非空字符串值
        """
        for value in values:
            text = str(value or "").strip()
            if text:
                return text
        return ""

    @classmethod
    def _build_vendor_store_options_from_store_configs(
        cls, store_configs: list[TicketLogPullStoreConfig]
    ) -> list[TicketLogPullVendorOptionModel]:
        """
        按门店配置表聚合商家与门店联动选项。
        :param store_configs: 门店配置列表
        :return: 商家门店联动选项
        """
        vendor_map: dict[str, dict[str, Any]] = {}
        vendor_order: list[str] = []
        for store_config in store_configs:
            vendor_code = str(store_config.vender_no or "").strip()
            vendor_id = cls._extract_option_id(store_config.vender_no, store_config.id)
            if vendor_id is None:
                continue
            vendor_key = vendor_code or str(vendor_id)
            vendor_entry = vendor_map.get(vendor_key)
            if not vendor_entry:
                vendor_entry = {
                    "vendor_id": vendor_id,
                    "vendor_code": vendor_code or None,
                    "vendor_name": vendor_code or str(store_config.group_no or "").strip() or "未命名商家",
                    "stores": [],
                    "store_ids": set(),
                }
                vendor_map[vendor_key] = vendor_entry
                vendor_order.append(vendor_key)

            store_value = cls._normalize_store_option_value(store_config.org_no, store_config.sap_org_no)
            if not store_value:
                continue
            store_key = store_value
            if store_key in vendor_entry["store_ids"]:
                continue
            vendor_entry["store_ids"].add(store_key)
            store_code = str(store_config.org_no or "").strip() or None
            sap_org_no = str(store_config.sap_org_no or "").strip() or None
            store_name = str(store_config.org_name or "").strip() or store_code or sap_org_no or str(
                store_config.vender_no or ""
            ).strip() or store_value
            vendor_entry["stores"].append(
                TicketLogPullStoreOptionModel(
                    store_id=store_value,
                    store_code=store_code,
                    sap_org_no=sap_org_no,
                    store_name=store_name,
                )
            )

        vendors: list[TicketLogPullVendorOptionModel] = []
        for vendor_key in vendor_order:
            vendor_entry = vendor_map[vendor_key]
            if not vendor_entry["stores"]:
                continue
            vendors.append(
                TicketLogPullVendorOptionModel(
                    vendor_id=int(vendor_entry["vendor_id"]),
                    vendor_code=vendor_entry["vendor_code"],
                    vendor_name=vendor_entry["vendor_name"],
                    stores=vendor_entry["stores"],
                )
            )
        return vendors

    @classmethod
    def get_vendor_store_options_services(
        cls, query_db: Session, vendor_id: int | None = None
    ) -> TicketLogPullVendorStoreOptionsModel:
        """
        获取日志拉取页面使用的商家/门店联动选项。
        :param query_db: 数据库会话
        :param vendor_id: 可选商家ID，传入后只返回该商家对应的门店
        :return: 商家/门店联动配置
        """
        store_configs = TicketLogPullDao.list_all_store_configs(query_db)
        if vendor_id is not None:
            resolved_vendor_id = int(vendor_id)
            store_configs = [
                store_config
                for store_config in store_configs
                if cls._extract_option_id(store_config.vender_no, store_config.id) == resolved_vendor_id
            ]
        cls.ensure_param_config_rows(query_db)
        config_row = TicketLogPullDao.get_param_example_config_row(query_db)
        raw_examples = cls._json_loads(getattr(config_row, "config_value", None), [])
        vendors = cls._build_vendor_store_options_from_store_configs(store_configs)
        return TicketLogPullVendorStoreOptionsModel(
            vendors=vendors,
            parameter_examples=cls._normalize_parameter_examples(raw_examples),
        )

    @classmethod
    def build_store_config_import_template(cls) -> bytes:
        """
        生成门店配置导入模板。
        :return: Excel 文件字节
        """
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "门店配置导入模板"
        for index, header in enumerate(cls.STORE_IMPORT_HEADERS, 1):
            cell = sheet.cell(row=1, column=index, value=header)
            sheet.column_dimensions[cell.column_letter].width = 18
            sheet.cell(row=2, column=index, value=cls.STORE_IMPORT_SAMPLE.get(header, ""))
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @classmethod
    def get_store_config_list_services(
        cls, query_db: Session, query: TicketLogPullStoreConfigQueryModel
    ):
        """
        查询门店配置列表。
        :param query_db: 数据库会话
        :param query: 查询参数
        :return: 分页结果
        """
        result = TicketLogPullDao.list_store_configs(query_db, query)
        if query.is_page:
            result.rows = [CamelCaseUtil.transform_result(row) for row in result.rows]
            return result
        return [CamelCaseUtil.transform_result(row) for row in result]

    @classmethod
    def import_store_config_services(
        cls,
        query_db: Session,
        file_content: bytes,
        import_mode: str,
        current_user: CurrentUserModel,
    ) -> CrudResponseModel:
        """
        导入门店配置。
        :param query_db: 数据库会话
        :param file_content: Excel 文件内容
        :param import_mode: 导入方式，incremental 或 overwrite
        :param current_user: 当前登录用户
        :return: 导入结果
        """
        workbook = load_workbook(filename=BytesIO(file_content), data_only=True)
        sheet = workbook.active
        header_map = cls._build_store_header_map([cell.value for cell in sheet[1]])
        rows: list[tuple[int, dict[str, Any]]] = []
        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            row_data = cls._read_store_row(row, header_map)
            if not any(value not in (None, "") for value in row_data.values()):
                continue
            rows.append((row_index, row_data))

        normalized_mode = str(import_mode or "incremental").strip().lower()
        if normalized_mode not in {"incremental", "overwrite"}:
            return CrudResponseModel(is_success=False, message="导入方式仅支持 incremental 或 overwrite")

        try:
            existing_store_map: dict[tuple[str, str, str], TicketLogPullStoreConfig] = {}
            if normalized_mode == "overwrite":
                deleted_count = TicketLogPullDao.delete_all_store_configs(query_db)
                cls._log_chain_step(
                    query_db,
                    ticket_id=None,
                    record_id=None,
                    step="store-import",
                    status="overwrite",
                    reason=f"覆盖导入前清空旧数据 {deleted_count} 条",
                    detail={"deletedCount": deleted_count},
                )
            else:
                # 先把已有配置放入内存，避免逐行查库和重复 flush。
                for store in TicketLogPullDao.list_all_store_configs(query_db):
                    match_key = cls._build_store_config_match_key(store)
                    if any(match_key):
                        existing_store_map[match_key] = store

            summary = {
                "totalRows": len(rows),
                "insertedCount": 0,
                "updatedCount": 0,
                "failedRows": [],
                "importMode": normalized_mode,
            }
            now = datetime.now()
            for row_index, row in rows:
                try:
                    store = cls._build_store_config_entity(row, now)
                    match_key = cls._build_store_config_match_key(store)
                    if not any(match_key):
                        raise ValueError("vender_no/org_no/sap_org_no 至少需要填写一个")

                    existing = existing_store_map.get(match_key)
                    if existing:
                        for field in (
                            "group_no",
                            "vender_no",
                            "region_no",
                            "org_no",
                            "org_name",
                            "sap_org_no",
                            "platform_no",
                            "parent_org_no",
                            "perm_node_id",
                            "org_type",
                            "company_no",
                            "city_no",
                            "biz_type_no",
                            "status",
                            "created",
                            "modifid",
                            "open_date",
                            "language_desc",
                        ):
                            setattr(existing, field, getattr(store, field))
                        saved = existing
                        summary["updatedCount"] += 1
                    else:
                        query_db.add(store)
                        existing_store_map[match_key] = store
                        saved = store
                        summary["insertedCount"] += 1

                    cls._log_chain_step(
                        query_db,
                        ticket_id=None,
                        record_id=int(saved.id) if getattr(saved, "id", None) else None,
                        step="store-import",
                        status="success",
                        reason="覆盖保存完成" if existing else "新增保存完成",
                        detail={
                            "row": row_index,
                            "venderNo": saved.vender_no,
                            "orgNo": saved.org_no,
                            "sapOrgNo": saved.sap_org_no,
                            "importMode": normalized_mode,
                        },
                    )
                except Exception as exc:
                    summary["failedRows"].append({"row": row_index, "reason": str(exc)})
                    cls._log_chain_step(
                        query_db,
                        ticket_id=None,
                        record_id=None,
                        step="store-import",
                        status="failed",
                        reason=str(exc),
                        detail={"row": row_index},
                    )
            query_db.commit()
            return CrudResponseModel(is_success=True, message="门店配置导入完成", result=summary)
        except Exception:
            query_db.rollback()
            raise

    @classmethod
    def get_project_vendor_map_list_services(cls, query_db: Session, query: TicketLogPullProjectVendorMapQueryModel):
        """
        查询项目商家映射列表。
        :param query_db: 数据库会话
        :param query: 查询参数
        :return: 分页结果
        """
        result = TicketLogPullDao.list_project_vendor_maps(query_db, query)
        if query.is_page:
            result.rows = [CamelCaseUtil.transform_result(row) for row in result.rows]
            return result
        return [CamelCaseUtil.transform_result(row) for row in result]

    @classmethod
    def get_project_vendor_map_options_services(cls, query_db: Session) -> list[dict[str, Any]]:
        """
        获取全部项目商家映射选项。
        :param query_db: 数据库会话
        :return: 映射列表
        """
        return [CamelCaseUtil.transform_result(row) for row in TicketLogPullDao.list_all_project_vendor_maps(query_db)]

    @classmethod
    def get_project_vendor_map_by_project_services(
        cls, query_db: Session, project_id: int
    ) -> TicketLogPullProjectVendorMapModel | None:
        """
        根据项目ID获取商家映射。
        :param query_db: 数据库会话
        :param project_id: 项目ID
        :return: 映射信息
        """
        row = TicketLogPullDao.get_project_vendor_map_by_project_id(query_db, project_id)
        return TicketLogPullProjectVendorMapModel.model_validate(row) if row else None

    @classmethod
    def save_project_vendor_map_services(
        cls,
        query_db: Session,
        payload: TicketLogPullProjectVendorMapUpsertModel,
        current_user: CurrentUserModel,
    ) -> CrudResponseModel:
        """
        保存项目商家映射。
        :param query_db: 数据库会话
        :param payload: 保存参数
        :param current_user: 当前登录用户
        :return: 保存结果
        """
        now = datetime.now()
        try:
            saved = TicketLogPullDao.save_project_vendor_map(
                query_db,
                TicketLogPullProjectVendorMap(
                    project_id=int(payload.project_id),
                    project_name=str(payload.project_name or "").strip(),
                    vender_no=str(payload.vender_no or "").strip(),
                    created=now,
                    modifid=now,
                ),
            )
            query_db.commit()
            cls._log_chain_step(
                query_db,
                ticket_id=None,
                record_id=int(saved.id),
                step="project-vendor-map",
                status="saved",
                reason="项目商家映射已保存",
                detail={
                    "projectId": saved.project_id,
                    "venderNo": saved.vender_no,
                    "user": cls._user_name(current_user),
                },
            )
            return CrudResponseModel(
                is_success=True,
                message="项目商家映射已保存",
                result=CamelCaseUtil.transform_result(saved),
            )
        except Exception:
            query_db.rollback()
            raise

    @classmethod
    def _build_store_config_entity(cls, row: dict[str, Any], now: datetime) -> TicketLogPullStoreConfig:
        """
        根据导入行构建门店配置实体。
        :param row: 导入行数据
        :param now: 当前时间
        :return: 门店配置实体
        """
        return TicketLogPullStoreConfig(
            group_no=cls._cell_text(row.get("group_no")),
            vender_no=cls._cell_text(row.get("vender_no")),
            region_no=cls._cell_text(row.get("region_no")),
            org_no=cls._cell_text(row.get("org_no")) or None,
            org_name=cls._cell_text(row.get("org_name")) or None,
            sap_org_no=cls._cell_text(row.get("sap_org_no")) or None,
            platform_no=cls._cell_text(row.get("platform_no")),
            parent_org_no=cls._cell_text(row.get("parent_org_no")) or None,
            perm_node_id=cls._parse_import_int(row.get("perm_node_id"), 0),
            org_type=cls._parse_import_int(row.get("org_type"), 1),
            company_no=cls._cell_text(row.get("company_no")),
            city_no=cls._cell_text(row.get("city_no")),
            biz_type_no=cls._cell_text(row.get("biz_type_no")) or "1",
            status=cls._parse_import_int(row.get("status"), 1),
            created=cls._parse_import_datetime(row.get("created")) or now,
            modifid=cls._parse_import_datetime(row.get("modifid")) or now,
            open_date=cls._parse_import_date(row.get("open_date")) or date(1900, 1, 1),
            language_desc=cls._cell_text(row.get("language_desc")) or "zh_HK",
        )

    @classmethod
    def _build_store_config_match_key(cls, store: TicketLogPullStoreConfig) -> tuple[str, str, str]:
        """
        构建门店配置去重键。
        :param store: 门店配置实体
        :return: 去重键值元组
        """
        # 唯一键由 vender_no + org_no + sap_org_no 共同组成，必须三者同时一致才认为重复。
        return (
            cls._cell_text(getattr(store, "vender_no", "")),
            cls._cell_text(getattr(store, "org_no", "")),
            cls._cell_text(getattr(store, "sap_org_no", "")),
        )

    @classmethod
    def _build_store_header_map(cls, headers: list[Any]) -> dict[str, int]:
        """
        构建门店配置导入表头映射。
        :param headers: 首行表头
        :return: 字段映射
        """
        alias_map = {
            "group_no": ["集团编号", "group_no", "groupNo"],
            "vender_no": ["商户编号", "vender_no", "venderNo", "vendor_no", "vendorNo"],
            "region_no": ["区域编号", "region_no", "regionNo"],
            "org_no": ["机构编号", "org_no", "orgNo"],
            "org_name": ["机构名称", "org_name", "orgName"],
            "sap_org_no": ["SAP机构编号", "sap_org_no", "sapOrgNo"],
            "platform_no": ["会员渠道编号", "platform_no", "platformNo"],
            "parent_org_no": ["上级机构编号", "parent_org_no", "parentOrgNo"],
            "perm_node_id": ["权限树节点ID", "perm_node_id", "permNodeId"],
            "org_type": ["机构类型", "org_type", "orgType"],
            "company_no": ["所属公司代码", "company_no", "companyNo"],
            "city_no": ["城市编号", "city_no", "cityNo"],
            "biz_type_no": ["业态编号", "biz_type_no", "bizTypeNo"],
            "status": ["状态", "status"],
            "created": ["创建时间", "created"],
            "modifid": ["修改时间", "modifid"],
            "open_date": ["开业日期", "open_date", "openDate"],
            "language_desc": ["默认语言", "language_desc", "languageDesc"],
        }
        result: dict[str, int] = {}
        for index, header in enumerate(headers):
            header_text = cls._cell_text(header)
            for field_name, aliases in alias_map.items():
                if header_text in aliases and field_name not in result:
                    result[field_name] = index
        return result

    @classmethod
    def _read_store_row(cls, row: tuple[Any, ...], header_map: dict[str, int]) -> dict[str, Any]:
        """
        读取门店配置导入行。
        :param row: 行数据
        :param header_map: 表头映射
        :return: 字段字典
        """
        data: dict[str, Any] = {}
        for field_name, index in header_map.items():
            data[field_name] = row[index] if index < len(row) else None
        return data

    @classmethod
    def ensure_param_config_rows(cls, query_db: Session) -> None:
        """
        初始化日志拉取相关系统参数，确保参数配置管理中存在默认项。
        :param query_db: 数据库会话
        :return: 无
        """
        if not TicketLogPullDao.get_storage_config_row(query_db):
            payload = cls._default_storage_config()
            payload.pop("effectiveLocalDirectory", None)
            TicketLogPullDao.save_storage_config_row(
                query_db,
                config_value=cls._json_dumps(payload),
                user_name="system",
            )
        if not TicketLogPullDao.get_external_config_row(query_db):
            TicketLogPullDao.save_external_config_row(
                query_db,
                config_value=cls._json_dumps(cls._default_external_config()),
                user_name="system",
            )
        if not TicketLogPullDao.get_param_example_config_row(query_db):
            TicketLogPullDao.save_config_row(
                query_db,
                config_key=TicketLogPullDao.PARAM_EXAMPLE_CONFIG_KEY,
                config_name=TicketLogPullDao.PARAM_EXAMPLE_CONFIG_NAME,
                config_value=cls._json_dumps(cls._default_parameter_examples()),
                user_name="system",
                remark="工单日志拉取 modifyTime/path 参数示例配置",
            )

    @classmethod
    def _build_external_request_headers(cls, config: dict[str, Any]) -> dict[str, str]:
        """
        构建请求外部日志平台时使用的请求头。
        :param config: 外部接口配置
        :return: 请求头字典
        """
        headers = config.get("headers") if isinstance(config.get("headers"), dict) else {}
        return {
            key: str(value).strip()
            for key, value in headers.items()
            if str(value or "").strip()
        }

    @classmethod
    def get_storage_config_services(cls, query_db: Session) -> TicketLogPullStorageConfigModel:
        """
        获取日志拉取存储配置。
        :param query_db: 数据库会话
        :return: 存储配置模型
        """
        return TicketLogPullStorageConfigModel.model_validate(cls._get_storage_config_dict(query_db))

    @classmethod
    def save_storage_config_services(
        cls, query_db: Session, config_model: TicketLogPullStorageConfigModel, current_user: CurrentUserModel
    ) -> CrudResponseModel:
        """
        保存日志拉取存储配置。
        :param query_db: 数据库会话
        :param config_model: 存储配置模型
        :param current_user: 当前登录用户
        :return: 保存结果
        """
        payload = cls._normalize_storage_config(config_model.model_dump(by_alias=True))
        payload.pop("effectiveLocalDirectory", None)
        try:
            TicketLogPullDao.save_storage_config_row(
                query_db,
                config_value=cls._json_dumps(payload),
                user_name=cls._user_name(current_user),
            )
            query_db.commit()
            return CrudResponseModel(is_success=True, message="日志拉取存储配置已保存")
        except Exception:
            query_db.rollback()
            raise

    @classmethod
    def retry_log_pull_services(
        cls, query_db: Session, record_id: int, current_user: CurrentUserModel
    ) -> CrudResponseModel:
        """
        基于原记录重新提交日志拉取任务。
        :param query_db: 数据库会话
        :param record_id: 原日志拉取记录ID
        :param current_user: 当前登录用户
        :return: 重新提交结果
        """
        record = TicketLogPullDao.get_record_by_id(query_db, record_id)
        if not record:
            cls._log_chain_step(
                query_db,
                ticket_id=None,
                record_id=record_id,
                step="retry-log-pull",
                status="skipped",
                reason="日志拉取记录不存在",
            )
            return CrudResponseModel(is_success=False, message="日志拉取记录不存在")
        if record.status in cls.ACTIVE_STATUSES:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="retry-log-pull",
                status="skipped",
                reason="当前日志拉取任务仍在执行中",
            )
            return CrudResponseModel(is_success=False, message="当前日志拉取任务仍在执行中，暂不能重新拉取")
        payload = cls._build_retry_payload(record)
        if not payload:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="retry-log-pull",
                status="skipped",
                reason="当前记录缺少可重新拉取的原始参数",
            )
            return CrudResponseModel(is_success=False, message="当前记录缺少可重新拉取的原始参数")
        cls._log_chain_step(
            query_db,
            ticket_id=record.ticket_id,
            record_id=record.id,
            step="retry-log-pull",
            status="requested",
            reason="重新拉取已重新提交",
        )
        return cls.create_log_pull_services(query_db, record.ticket_id, payload, current_user)

    @classmethod
    def _extract_version_key_from_text(cls, text: str | None) -> str:
        """
        从日志文本中提取版本号。

        :param text: 日志文本。
        :return: 版本号，失败返回空字符串。
        """
        if not text:
            return ""
        match = cls.VERSION_PATTERN.search(text)
        if not match:
            return ""
        return str(match.group(1) or "").strip()

    @classmethod
    def _update_ticket_version_key(cls, query_db: Session, ticket_id: int, version_key: str) -> bool:
        """
        回写工单版本号到扩展字段。

        :param query_db: 数据库会话。
        :param ticket_id: 工单ID。
        :param version_key: 版本号。
        :return: 是否发生更新。
        """
        ticket = TicketDao.get_ticket_by_id(query_db, ticket_id)
        if not ticket:
            return False
        normalized_version_key = str(version_key or "").strip()
        if not normalized_version_key:
            return False
        extra_data = dict(ticket.extra_data or {}) if isinstance(ticket.extra_data, dict) else {}
        current_extra_version = str(extra_data.get("version_key") or "").strip()
        if current_extra_version == normalized_version_key:
            return False
        extra_data["version_key"] = normalized_version_key
        TicketDao.update_ticket(
            query_db,
            ticket_id,
            {
                "extra_data": extra_data,
                "update_by": "system",
                "update_time": datetime.now(),
            },
        )
        query_db.commit()
        return True

    @classmethod
    def _ensure_ticket_version_key_from_log(cls, query_db: Session, ticket_id: int, record_id: int) -> str:
        """
        日志拉取成功后确保工单具备版本号，缺失时从日志正文提取并回填。
        :param query_db: 数据库会话
        :param ticket_id: 工单ID
        :param record_id: 日志拉取记录ID
        :return: 可用版本号，未命中返回空字符串
        """
        ticket = TicketDao.get_ticket_by_id(query_db, ticket_id)
        if not ticket:
            return ""
        extra_data = ticket.extra_data if isinstance(ticket.extra_data, dict) else {}
        version_key = str(extra_data.get("version_key") or "").strip()
        if version_key:
            return version_key

        log_text = ""
        try:
            log_content_model = cls.get_log_pull_content_services(query_db, record_id)
            log_text = log_content_model.text if log_content_model else ""
        except Exception as exc:
            logger.warning(f"日志拉取记录[{record_id}] 提取版本号前读取日志失败: {exc}")
        version_key = cls._extract_version_key_from_text(log_text)
        if not version_key:
            cls._log_chain_step(
                query_db,
                ticket_id=ticket_id,
                record_id=record_id,
                step="version-backfill",
                status="skipped",
                reason="未从日志中提取到版本号",
            )
            return ""
        try:
            updated = cls._update_ticket_version_key(query_db, ticket_id, version_key)
            cls._log_chain_step(
                query_db,
                ticket_id=ticket_id,
                record_id=record_id,
                step="version-backfill",
                status="success" if updated else "skipped",
                reason="已从日志提取并回填版本号" if updated else "工单版本号已是最新值",
                detail={"versionKey": version_key},
            )
        except Exception as exc:
            query_db.rollback()
            logger.warning(f"日志拉取记录[{record_id}] 回填工单版本号失败: {exc}")
            cls._log_chain_step(
                query_db,
                ticket_id=ticket_id,
                record_id=record_id,
                step="version-backfill",
                status="failed",
                reason=str(exc),
            )
            return ""
        return version_key

    @classmethod
    def _notify_automation(
        cls,
        query_db: Session,
        ticket_id: int | None,
        *,
        status: str,
        message: str,
        detail: str | None = None,
        notify_config: dict[str, Any] | None = None,
    ) -> None:
        """
        按工单自动化配置发送通知。

        :param query_db: 数据库会话。
        :param ticket_id: 工单ID。
        :param status: 状态，success 或 failed。
        :param message: 简要说明。
        :param detail: 额外说明。
        :return: 无。
        """
        if not ticket_id:
            return
        ticket = TicketDao.get_ticket_by_id(query_db, ticket_id)
        if not ticket:
            return
        TicketNotifyService.send_ticket_notification(
            query_db,
            ticket,
            title="工单自动化通知",
            status=status,
            message=message,
            detail=detail,
            notify_config=notify_config,
        )

    @staticmethod
    def _extract_record_notify_config(record: TicketLogPullRecord | None) -> dict[str, Any]:
        """
        从日志拉取记录原始命令中提取通知配置。
        :param record: 日志拉取记录
        :return: 通知配置字典，未配置时返回空字典
        """
        command_content = getattr(record, "command_content", None)
        if not isinstance(command_content, dict):
            return {}
        notify_config = command_content.get("notifyConfig") or command_content.get("notify_config")
        if isinstance(notify_config, dict):
            return notify_config
        automation = command_content.get("_automation")
        if isinstance(automation, dict):
            nested_notify = automation.get("notifyConfig") or automation.get("notify_config")
            if isinstance(nested_notify, dict):
                return nested_notify
        return {}

    @classmethod
    def _notify_log_pull_record(
        cls,
        query_db: Session,
        record: TicketLogPullRecord | None,
        *,
        status: str,
        message: str,
        detail: str | None = None,
    ) -> None:
        """
        按日志拉取记录中的通知配置发送拉取结果通知。
        :param query_db: 数据库会话
        :param record: 日志拉取记录
        :param status: 通知状态，success 或 failed
        :param message: 通知说明
        :param detail: 通知详情
        :return: 无
        """
        if not record:
            return
        notify_config = cls._extract_record_notify_config(record)
        if not notify_config:
            return
        ticket = TicketDao.get_ticket_by_id(query_db, record.ticket_id) if record.ticket_id else None
        TicketNotifyService.send_ticket_notification(
            query_db,
            ticket,
            title="日志拉取通知",
            status=status,
            message=message,
            detail=detail,
            notify_config=notify_config,
        )

    @classmethod
    def redownload_log_pull_services(
        cls, query_db: Session, record_id: int, current_user: CurrentUserModel
    ) -> CrudResponseModel:
        """
        重新下载日志压缩包并恢复到原存储位置。
        :param query_db: 数据库会话
        :param record_id: 日志拉取记录ID
        :param current_user: 当前登录用户
        :return: 重新下载结果
        """
        record = TicketLogPullDao.get_record_by_id(query_db, record_id)
        if not record:
            cls._log_chain_step(
                query_db,
                ticket_id=None,
                record_id=record_id,
                step="redownload-log-pull",
                status="skipped",
                reason="日志拉取记录不存在",
            )
            return CrudResponseModel(is_success=False, message="日志拉取记录不存在")
        if record.status in cls.ACTIVE_STATUSES:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="redownload-log-pull",
                status="skipped",
                reason="当前日志拉取任务仍在执行中",
            )
            return CrudResponseModel(is_success=False, message="当前日志拉取任务仍在执行中，暂不能重新下载")
        if not str(record.command_result_url or "").strip():
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="redownload-log-pull",
                status="skipped",
                reason="当前记录缺少原始压缩包地址",
            )
            return CrudResponseModel(is_success=False, message="当前记录缺少原始压缩包地址，无法重新下载")

        temp_file_path: Path | None = None
        try:
            temp_file_path, file_size = cls._download_archive(record, query_db)
            storage_path = cls._restore_archive_storage(record, temp_file_path, query_db)
            now = datetime.now()
            TicketLogPullDao.update_record(
                query_db,
                record.id,
                {
                    "download_file_name": temp_file_path.name,
                    "download_file_size": file_size,
                    "storage_path": storage_path,
                    "update_by": cls._user_name(current_user),
                    "update_time": now,
                },
            )
            cls._add_ticket_event(
                query_db,
                ticket_id=record.ticket_id,
                operator_id=cls._user_id(current_user),
                operator_name=cls._user_name(current_user),
                content="重新下载日志压缩包",
                event_data={
                    "record_id": record.id,
                    "storage_mode": record.storage_mode,
                    "storage_path": storage_path,
                    "download_file_name": temp_file_path.name,
                    "download_file_size": file_size,
                },
            )
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="redownload-log-pull",
                status="success",
                reason="日志压缩包已重新下载",
                detail={"storagePath": storage_path, "fileSize": file_size},
            )
            query_db.commit()
            return CrudResponseModel(is_success=True, message="日志压缩包已重新下载")
        except Exception:
            query_db.rollback()
            raise
        finally:
            if temp_file_path and temp_file_path.exists():
                try:
                    temp_file_path.unlink()
                except Exception:
                    logger.warning("删除重新下载产生的临时文件失败: %s", temp_file_path)

    @classmethod
    def delete_log_pull_services(
        cls, query_db: Session, record_id: int, current_user: CurrentUserModel
    ) -> CrudResponseModel:
        """
        删除日志拉取记录并同步清理对应文件。
        :param query_db: 数据库会话
        :param record_id: 日志拉取记录ID
        :param current_user: 当前登录用户
        :return: 删除结果
        """
        record = TicketLogPullDao.get_record_by_id(query_db, record_id)
        if not record:
            cls._log_chain_step(
                query_db,
                ticket_id=None,
                record_id=record_id,
                step="delete-log-pull",
                status="skipped",
                reason="日志拉取记录不存在",
            )
            return CrudResponseModel(is_success=False, message="日志拉取记录不存在")
        if record.status in cls.ACTIVE_STATUSES:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="delete-log-pull",
                status="skipped",
                reason="当前日志拉取任务仍在执行中",
            )
            return CrudResponseModel(is_success=False, message="当前日志拉取任务仍在执行中，暂不能删除")

        try:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="delete-log-pull",
                status="requested",
                reason="准备删除日志拉取记录并清理文件",
            )
            cls._delete_record_storage(record, query_db)
            query_db.delete(record)
            cls._add_ticket_event(
                query_db,
                ticket_id=record.ticket_id,
                operator_id=cls._user_id(current_user),
                operator_name=cls._user_name(current_user),
                content="删除日志拉取记录",
                event_data={
                    "record_id": record.id,
                    "storage_mode": record.storage_mode,
                    "storage_path": record.storage_path,
                },
            )
            query_db.commit()
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="delete-log-pull",
                status="success",
                reason="日志拉取记录已删除",
            )
            return CrudResponseModel(is_success=True, message="日志拉取记录已删除")
        except Exception:
            query_db.rollback()
            raise

    @classmethod
    def reextract_log_pull_services(
        cls,
        query_db: Session,
        record_id: int,
        query: TicketLogPullContentQueryModel,
        current_user: CurrentUserModel,
    ) -> CrudResponseModel:
        """
        按新的时间范围重新截取并更新日志入库内容。
        :param query_db: 数据库会话
        :param record_id: 日志拉取记录ID
        :param query: 日志查看时间范围参数
        :param current_user: 当前登录用户
        :return: 重新截取结果
        """
        record = TicketLogPullDao.get_record_by_id(query_db, record_id)
        if not record:
            cls._log_chain_step(
                query_db,
                ticket_id=None,
                record_id=record_id,
                step="reextract-log-pull",
                status="skipped",
                reason="日志拉取记录不存在",
            )
            return CrudResponseModel(is_success=False, message="日志拉取记录不存在")
        if record.status in cls.ACTIVE_STATUSES:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="reextract-log-pull",
                status="skipped",
                reason="当前日志拉取任务仍在执行中",
            )
            return CrudResponseModel(is_success=False, message="当前日志拉取任务仍在执行中，暂不能重新截取")

        begin_time, end_time = cls._resolve_view_log_time_range(query)
        if not begin_time or not end_time:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="reextract-log-pull",
                status="skipped",
                reason="重新截取时日志时间范围必填",
            )
            return CrudResponseModel(is_success=False, message="重新截取时日志时间范围必填")
        if begin_time > end_time:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="reextract-log-pull",
                status="skipped",
                reason="开始时间不能晚于结束时间",
            )
            return CrudResponseModel(is_success=False, message="重新截取时开始时间不能晚于结束时间")

        archive_path, should_cleanup = cls._resolve_archive_source_for_view(record, query_db)
        if not archive_path:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="reextract-log-pull",
                status="skipped",
                reason="当前记录没有可用压缩包",
            )
            return CrudResponseModel(is_success=False, message="当前记录没有可用的压缩包文件，无法重新截取")

        try:
            content_result = cls._extract_archive_content(
                record,
                archive_path,
                query_db,
                begin_time=begin_time,
                end_time=end_time,
            )
            now = datetime.now()
            TicketLogPullDao.update_record(
                query_db,
                record.id,
                {
                    "log_begin_time": begin_time,
                    "log_end_time": end_time,
                    "status": TicketLogPullStatus.SUCCESS.value,
                    "status_desc": "日志已按当前时间范围重新截取",
                    "is_error": False,
                    "error_message": None,
                    "exception_detail": None,
                    "archive_entry_count": content_result["archive_entry_count"],
                    "matched_entry_count": content_result["matched_entry_count"],
                    "content_char_count": content_result["content_char_count"],
                    "content_truncated": content_result["content_truncated"],
                    "compressed_content": content_result["compressed_content"],
                    "content_summary": content_result["content_summary"],
                    "finished_at": now,
                    "update_by": cls._user_name(current_user),
                    "update_time": now,
                },
            )
            cls._add_ticket_event(
                query_db,
                ticket_id=record.ticket_id,
                operator_id=cls._user_id(current_user),
                operator_name=cls._user_name(current_user),
                content="重新截取日志内容",
                event_data={
                    "record_id": record.id,
                    "view_begin_time": begin_time,
                    "view_end_time": end_time,
                    "matched_entry_count": content_result["matched_entry_count"],
                    "archive_entry_count": content_result["archive_entry_count"],
                },
            )
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="reextract-log-pull",
                status="success",
                reason="重新截取完成",
                detail={
                    "matchedEntryCount": content_result["matched_entry_count"],
                    "archiveEntryCount": content_result["archive_entry_count"],
                },
            )
            query_db.commit()
            return CrudResponseModel(is_success=True, message="日志已按当前时间范围重新截取")
        except Exception:
            query_db.rollback()
            raise
        finally:
            if should_cleanup and archive_path:
                try:
                    archive_path.unlink(missing_ok=True)
                except Exception:
                    logger.warning("删除重新截取产生的临时文件失败: %s", archive_path)

    @classmethod
    def create_log_pull_services(
        cls, query_db: Session, ticket_id: int | None, payload: TicketLogPullCreateModel, current_user: CurrentUserModel
    ) -> CrudResponseModel:
        """
        创建日志拉取记录并触发后台执行。
        :param query_db: 数据库会话
        :param ticket_id: 工单ID，允许为空表示独立管理记录
        :param payload: 日志拉取参数
        :param current_user: 当前登录用户
        :return: 创建结果
        """
        resolved_ticket_id = ticket_id if ticket_id is not None else payload.ticket_id
        ticket = TicketDao.get_ticket_by_id(query_db, resolved_ticket_id) if resolved_ticket_id else None
        if resolved_ticket_id and not ticket:
            cls._log_chain_step(
                query_db,
                ticket_id=resolved_ticket_id,
                record_id=None,
                step="create-log-pull",
                status="skipped",
                reason="工单不存在",
            )
            return CrudResponseModel(is_success=False, message="工单不存在")

        storage_config = cls._get_storage_config_dict(query_db)
        storage_mode = str(payload.storage_mode or storage_config.get("mode") or "local").strip().lower()
        if storage_mode not in {"local", "ftp"}:
            storage_mode = "local"
        now = datetime.now()
        command_content = payload.model_dump(by_alias=True)
        log_begin_time, log_end_time = cls._resolve_log_time_range(payload)
        has_explicit_range = cls._has_explicit_log_time_range(payload)
        if has_explicit_range and (not log_begin_time or not log_end_time):
            cls._log_chain_step(
                query_db,
                ticket_id=resolved_ticket_id,
                record_id=None,
                step="create-log-pull",
                status="skipped",
                reason="日志时间范围填写不完整",
            )
            return CrudResponseModel(is_success=False, message="日志时间范围填写不完整")
        if log_begin_time and log_end_time and log_begin_time > log_end_time:
            cls._log_chain_step(
                query_db,
                ticket_id=resolved_ticket_id,
                record_id=None,
                step="create-log-pull",
                status="skipped",
                reason="日志开始时间不能晚于结束时间",
            )
            return CrudResponseModel(is_success=False, message="日志开始时间不能晚于结束时间")

        try:
            record = TicketLogPullDao.add_record(
                query_db,
                TicketLogPullRecord(
                    ticket_id=resolved_ticket_id,
                    vendor_id=payload.vendor_id,
                    store_id=payload.store_id,
                    pos_no=payload.pos_no,
                    command_type=1,
                    command_data_type=payload.command_data_type,
                    command_content=command_content,
                    log_begin_time=log_begin_time if has_explicit_range else None,
                    log_end_time=log_end_time if has_explicit_range else None,
                    storage_mode=storage_mode,
                    status=TicketLogPullStatus.CREATED.value,
                    status_desc="已创建，等待后台执行",
                    is_error=False,
                    create_by=cls._user_name(current_user),
                    update_by=cls._user_name(current_user),
                    create_time=now,
                    update_time=now,
                ),
            )
            cls._add_ticket_event(
                query_db,
                ticket_id=resolved_ticket_id,
                operator_id=cls._user_id(current_user),
                operator_name=cls._user_name(current_user),
                content="提交日志拉取申请",
                event_data={
                    "record_id": record.id,
                    "vendor_id": payload.vendor_id,
                    "store_id": payload.store_id,
                    "pos_no": payload.pos_no,
                    "command_data_type": payload.command_data_type,
                    "storage_mode": storage_mode,
                    "log_begin_time": log_begin_time,
                    "log_end_time": log_end_time,
                    "auto_ai_enabled": bool(payload.auto_ai_enabled),
                    "ai_agent_code": str(payload.ai_agent_code or "").strip() or None,
                    "ai_provider_code": str(payload.ai_provider_code or "").strip() or None,
                },
            )
            cls._log_chain_step(
                query_db,
                ticket_id=resolved_ticket_id,
                record_id=record.id,
                step="create-log-pull",
                status="created",
                reason="日志拉取记录已创建，等待后台执行",
                detail={
                    "vendorId": payload.vendor_id,
                    "storeId": payload.store_id,
                    "posNo": payload.pos_no,
                    "commandDataType": payload.command_data_type,
                },
            )
            query_db.commit()
            cls.queue_record(record.id)
            cls._log_chain_step(
                query_db,
                ticket_id=resolved_ticket_id,
                record_id=record.id,
                step="create-log-pull",
                status="queued",
                reason="后台任务已入队",
            )
            return CrudResponseModel(
                is_success=True,
                message="日志拉取任务已提交",
                result=CamelCaseUtil.transform_result(record),
            )
        except Exception:
            query_db.rollback()
            raise

    @classmethod
    def list_log_pull_records_services(cls, query_db: Session, query: TicketLogPullQueryModel):
        """
        查询工单日志拉取记录列表。
        :param query_db: 数据库会话
        :param query: 查询参数
        :return: 分页结果
        """
        result = TicketLogPullDao.list_ticket_records(query_db, query)
        if query.is_page:
            result.rows = [cls._to_record_list_item(row) for row in result.rows]
            return result
        return [cls._to_record_list_item(row) for row in result]

    @classmethod
    def list_log_pull_management_records_services(
        cls, query_db: Session, query: TicketLogPullQueryModel
    ):
        """
        查询日志拉取管理页记录。
        :param query_db: 数据库会话
        :param query: 查询参数
        :return: 分页结果或列表
        """
        result = TicketLogPullDao.list_log_pull_records(query_db, query)
        if query.is_page:
            result.rows = cls._enrich_record_list_items(query_db, result.rows)
            return result
        return cls._enrich_record_list_items(query_db, result)

    @classmethod
    def download_log_pull_file_services(
        cls, query_db: Session, record_id: int, source: str = "auto"
    ) -> tuple[Path | None, bool, str | None]:
        """
        获取日志拉取记录对应的可下载文件。
        :param query_db: 数据库会话
        :param record_id: 日志拉取记录ID
        :param source: 下载来源，auto 优先本服务归档并回退原始地址，service 仅本服务归档，original 仅原始地址
        :return: 文件路径、是否需要清理临时文件、下载文件名
        """
        record = TicketLogPullDao.get_record_by_id(query_db, record_id)
        if not record:
            cls._log_chain_step(
                query_db,
                ticket_id=None,
                record_id=record_id,
                step="download-log-pull",
                status="skipped",
                reason="日志拉取记录不存在",
            )
            return None, False, None
        normalized_source = str(source or "auto").strip().lower()
        if normalized_source not in {"auto", "service", "original"}:
            normalized_source = "auto"
        archive_path, should_cleanup = cls._resolve_archive_source_for_download(
            record, query_db, source=normalized_source
        )
        if not archive_path:
            cls._log_chain_step(
                query_db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="download-log-pull",
                status="skipped",
                reason="没有可下载的归档文件",
                detail={"source": normalized_source},
            )
            return None, False, None
        cls._log_chain_step(
            query_db,
            ticket_id=record.ticket_id,
            record_id=record.id,
            step="download-log-pull",
            status="requested",
            reason="准备下载归档文件",
            detail={"source": normalized_source},
        )
        return archive_path, should_cleanup, cls._build_download_file_name(record, archive_path)

    @classmethod
    def get_log_pull_content_services(
        cls, query_db: Session, record_id: int, query: TicketLogPullContentQueryModel | None = None
    ) -> TicketLogPullContentModel | None:
        """
        获取日志拉取记录中的解压文本内容。
        :param query_db: 数据库会话
        :param record_id: 日志拉取记录ID
        :param query: 查看日志时的时间范围参数
        :return: 日志文本内容
        """
        record = TicketLogPullDao.get_record_by_id(query_db, record_id)
        if not record:
            return None
        text = cls._decompress_text(record.compressed_content)
        content_summary = record.content_summary
        view_begin_time = record.log_begin_time
        view_end_time = record.log_end_time
        view_source = "stored"
        content_char_count = record.content_char_count or len(text)
        content_truncated = bool(record.content_truncated)
        matched_entry_count = record.matched_entry_count or 0
        archive_entry_count = record.archive_entry_count or 0
        view_mode = str(getattr(query, "view_mode", "stored") or "stored").strip().lower()
        if view_mode == "archive":
            archive_path, should_cleanup = cls._resolve_archive_source_for_view(record, query_db)
            requested_range = cls._resolve_view_log_time_range(query) if query else (None, None)
            requested_begin_time, requested_end_time = requested_range
            view_begin_time = requested_begin_time or view_begin_time
            view_end_time = requested_end_time or view_end_time
            if archive_path:
                try:
                    content_result = cls._extract_archive_content(
                        record,
                        archive_path,
                        query_db,
                        begin_time=requested_begin_time,
                        end_time=requested_end_time,
                    )
                    text = content_result["compressed_content"]
                    content_summary = content_result["content_summary"]
                    content_char_count = content_result["content_char_count"]
                    content_truncated = content_result["content_truncated"]
                    matched_entry_count = content_result["matched_entry_count"]
                    archive_entry_count = content_result["archive_entry_count"]
                    view_source = "realtime"
                except Exception as exc:
                    logger.warning("实时截取日志失败，回退到已入库内容: %s", exc)
                    content_summary = cls._build_view_fallback_summary(
                        record.content_summary, "实时截取失败，已回退到入库内容"
                    )
                    view_source = "fallback"
            else:
                content_summary = cls._build_view_fallback_summary(
                    record.content_summary, "归档文件不可用，已返回入库内容"
                )
                view_source = "fallback"
            if should_cleanup and archive_path:
                try:
                    archive_path.unlink(missing_ok=True)
                except Exception:
                    pass
        return TicketLogPullContentModel(
            record_id=record.id,
            view_begin_time=view_begin_time,
            view_end_time=view_end_time,
            view_source=view_source,
            content_summary=content_summary,
            text=text,
            content_char_count=content_char_count,
            content_truncated=content_truncated,
            matched_entry_count=matched_entry_count,
            archive_entry_count=archive_entry_count,
            storage_path=record.storage_path,
            command_result_url=record.command_result_url,
        )

    @classmethod
    def get_latest_summary_map(cls, query_db: Session, ticket_ids: list[int]) -> dict[int, dict[str, Any]]:
        """
        查询多个工单的最新日志拉取摘要。
        :param query_db: 数据库会话
        :param ticket_ids: 工单ID列表
        :return: 以工单ID为键的摘要映射
        """
        latest_map = TicketLogPullDao.list_latest_records_by_ticket_ids(query_db, ticket_ids)
        return {
            ticket_id: TicketLogPullSummaryModel.model_validate(
                cls._to_record_summary(record)
            ).model_dump(by_alias=True)
            for ticket_id, record in latest_map.items()
        }

    @classmethod
    def get_latest_summary(cls, query_db: Session, ticket_id: int) -> dict[str, Any] | None:
        """
        查询单个工单的最新日志拉取摘要。
        :param query_db: 数据库会话
        :param ticket_id: 工单ID
        :return: 最新摘要
        """
        summary_map = cls.get_latest_summary_map(query_db, [ticket_id])
        return summary_map.get(ticket_id)

    @classmethod
    def queue_record(cls, record_id: int) -> None:
        """
        将日志拉取记录放入后台线程池执行。
        :param record_id: 记录ID
        :return: 无
        """
        with cls._executor_lock:
            if record_id in cls._active_record_ids:
                return
            cls._active_record_ids.add(record_id)
        cls._executor.submit(cls._run_record, record_id)

    @classmethod
    def resume_pending_records(cls) -> None:
        """
        服务启动后清理未完成的日志拉取记录。
        :return: 无
        """
        with SessionLocal() as db:
            records = TicketLogPullDao.list_recoverable_records(db, cls.ACTIVE_STATUSES)
            now = datetime.now()
            for record in records:
                cls._fail_record(
                    db,
                    record.id,
                    status=TicketLogPullStatus.FAILED.value,
                    status_desc="服务重启前任务未完成，已清理为失败",
                    error_message="服务重启前任务未完成，已清理为失败",
                )
                cls._add_ticket_event(
                    db,
                    ticket_id=record.ticket_id,
                    operator_id=None,
                    operator_name="system",
                    content="日志拉取任务因服务重启清理为失败",
                    event_data={"record_id": record.id, "status": record.status, "cleaned_at": now.isoformat()},
                )
            if records:
                db.commit()

    @classmethod
    def _run_record(cls, record_id: int) -> None:
        """
        在线程池中执行单条日志拉取记录。
        :param record_id: 记录ID
        :return: 无
        """
        try:
            with SessionLocal() as db:
                cls._process_record(db, record_id)
        except Exception as exc:
            logger.exception(exc)
        finally:
            with cls._executor_lock:
                cls._active_record_ids.discard(record_id)

    @classmethod
    def _process_record(cls, db: Session, record_id: int) -> None:
        """
        执行日志拉取全流程。
        :param db: 数据库会话
        :param record_id: 记录ID
        :return: 无
        """
        record = TicketLogPullDao.get_record_by_id(db, record_id)
        if not record or record.status == TicketLogPullStatus.SUCCESS.value:
            return

        temp_file_path: Path | None = None
        direct_download_ready = False
        try:
            has_log_time_range = bool(record.log_begin_time or record.log_end_time)
            if record.status == TicketLogPullStatus.CREATED.value:
                rows = cls._fetch_external_rows(db, record)
                matched_row = cls._match_external_row(record, rows)
                if matched_row:
                    external_status = int(matched_row.get("commandStatus") or 0)
                    command_result_url = cls._extract_command_result_url(matched_row.get("commandResult"))
                    if external_status == 2:
                        cls._fail_record(
                            db,
                            record_id,
                            status=TicketLogPullStatus.FAILED.value,
                            status_desc="外部平台日志拉取失败",
                            error_message=str(matched_row.get("errorMsg") or "外部平台返回失败"),
                        )
                        cls._add_ticket_event(
                            db,
                            ticket_id=record.ticket_id,
                            operator_id=None,
                            operator_name="system",
                            content="外部平台日志拉取失败",
                            event_data={"record_id": record.id, "error_message": matched_row.get("errorMsg")},
                        )
                        cls._notify_log_pull_record(
                            db,
                            record,
                            status="failed",
                            message="外部平台日志拉取失败",
                            detail=f"record_id={record.id}, error={matched_row.get('errorMsg')}",
                        )
                        db.commit()
                        return
                    if external_status == 1 and command_result_url:
                        now = datetime.now()
                        cls._update_status(
                            db,
                            record_id,
                            status=TicketLogPullStatus.DOWNLOADING.value,
                            status_desc="已匹配到外部平台结果，正在下载日志压缩包",
                            is_error=False,
                            update_by="system",
                            last_polled_at=now,
                            external_command_id=matched_row.get("id"),
                            external_serial_number=matched_row.get("serialNumber"),
                            external_command_status=external_status,
                            external_command_status_desc=matched_row.get("commandStatusDesc"),
                            command_result_url=command_result_url,
                            source_created_at=cls._parse_external_datetime(matched_row),
                        )
                        record = TicketLogPullDao.get_record_by_id(db, record_id)
                        if record is None:
                            return
                        direct_download_ready = True
                if record.status == TicketLogPullStatus.CREATED.value:
                    cls._update_status(
                        db,
                        record_id,
                        status=TicketLogPullStatus.SUBMITTING.value,
                        status_desc="正在提交外部日志拉取申请",
                        is_error=False,
                        update_by="system",
                    )
                    record = TicketLogPullDao.get_record_by_id(db, record_id)
                    if record is None:
                        return
                    cls._submit_external_request(db, record)
                    record = TicketLogPullDao.get_record_by_id(db, record_id)
                    if record is None:
                        return

            if record.status in {
                TicketLogPullStatus.SUBMITTING.value,
                TicketLogPullStatus.POLLING.value,
                TicketLogPullStatus.DOWNLOADING.value,
                TicketLogPullStatus.PROCESSING.value,
            }:
                if not direct_download_ready:
                    command_row = cls._poll_external_result(db, record)
                    record = TicketLogPullDao.get_record_by_id(db, record_id)
                    if record is None:
                        return
                    if not command_row:
                        return
                    if not record.command_result_url:
                        cls._fail_record(
                            db,
                            record_id,
                            status=TicketLogPullStatus.FAILED.value,
                            status_desc="外部平台未返回压缩包地址",
                            error_message="外部平台返回成功状态但缺少 commandResult.url",
                        )
                        cls._notify_log_pull_record(
                            db,
                            record,
                            status="failed",
                            message="外部平台未返回压缩包地址",
                            detail=f"record_id={record.id}",
                        )
                        return

                if record.status != TicketLogPullStatus.DOWNLOADING.value:
                    cls._update_status(
                        db,
                        record_id,
                        status=TicketLogPullStatus.DOWNLOADING.value,
                        status_desc="正在下载日志压缩包",
                        is_error=False,
                        update_by="system",
                    )
                record = TicketLogPullDao.get_record_by_id(db, record_id)
                if record is None:
                    return
                temp_file_path, file_size = cls._download_archive(record, db)
                storage_path = cls._store_archive(record, temp_file_path, db)
                update_kwargs: dict[str, Any] = {
                    "download_file_name": temp_file_path.name,
                    "download_file_size": file_size,
                    "storage_path": storage_path,
                }
                if has_log_time_range:
                    cls._update_status(
                        db,
                        record_id,
                        status=TicketLogPullStatus.PROCESSING.value,
                        status_desc="正在解析日志内容并压缩入库",
                        is_error=False,
                        update_by="system",
                        **update_kwargs,
                    )
                    record = TicketLogPullDao.get_record_by_id(db, record_id)
                    if record is None:
                        return
                    try:
                        content_result = cls._extract_archive_content(record, temp_file_path, db)
                    except TicketLogContentTooLargeError as exc:
                        cls._fail_record(
                            db,
                            record_id,
                            status=TicketLogPullStatus.FAILED.value,
                            status_desc="日志内容超出入库上限",
                            error_message=str(exc),
                        )
                        cls._notify_log_pull_record(
                            db,
                            record,
                            status="failed",
                            message="日志内容超出入库上限",
                            detail=f"record_id={record.id}, error={exc}",
                        )
                        return
                    success_desc = "日志拉取完成"
                    if record.command_data_type == TicketLogDataType.DB.value:
                        success_desc = "DB 拉取完成"
                    elif content_result["matched_entry_count"] == 0:
                        success_desc = "日志拉取完成，未匹配到时间范围内日志"
                    cls._update_status(
                        db,
                        record_id,
                        status=TicketLogPullStatus.SUCCESS.value,
                        status_desc=success_desc,
                        is_error=False,
                        update_by="system",
                        archive_entry_count=content_result["archive_entry_count"],
                        matched_entry_count=content_result["matched_entry_count"],
                        content_char_count=content_result["content_char_count"],
                        content_truncated=content_result["content_truncated"],
                        compressed_content=content_result["compressed_content"],
                        content_summary=content_result["content_summary"],
                        finished_at=datetime.now(),
                        **update_kwargs,
                    )
                else:
                    archive_entry_count = cls._count_archive_entries(temp_file_path)
                    cls._update_status(
                        db,
                        record_id,
                        status=TicketLogPullStatus.SUCCESS.value,
                        status_desc="日志已下载，未截取内容",
                        is_error=False,
                        update_by="system",
                        archive_entry_count=archive_entry_count,
                        matched_entry_count=0,
                        content_char_count=0,
                        content_truncated=False,
                        compressed_content=None,
                        content_summary="日志已下载完成，未截取入库，AI 分析将使用整包压缩文件。",
                        finished_at=datetime.now(),
                        **update_kwargs,
                    )
                cls._log_chain_step(
                    db,
                    ticket_id=record.ticket_id,
                    record_id=record.id,
                    step="archive-process",
                    status="success",
                    reason="归档与解析完成" if has_log_time_range else "仅归档完成",
                    detail={
                        "storagePath": storage_path,
                        "fileSize": file_size,
                        "hasLogTimeRange": has_log_time_range,
                    },
                )
                cls._add_ticket_event(
                    db,
                    ticket_id=record.ticket_id,
                    operator_id=None,
                    operator_name="system",
                    content="日志拉取完成" if has_log_time_range else "日志已下载，未截取内容",
                    event_data={
                        "record_id": record.id,
                        "storage_mode": record.storage_mode,
                        "storage_path": storage_path,
                        "matched_entry_count": content_result["matched_entry_count"] if has_log_time_range else 0,
                        "archive_entry_count": (
                            content_result["archive_entry_count"] if has_log_time_range else archive_entry_count
                        ),
                        "whole_archive": not has_log_time_range,
                    },
                )
                cls._notify_log_pull_record(
                    db,
                    record,
                    status="success",
                    message="日志拉取已完成" if has_log_time_range else "日志压缩包已下载，未切割入库",
                    detail=f"record_id={record.id}, storage_path={storage_path}",
                )
                db.commit()
                cls._trigger_auto_ai_analysis(db, record.id)
        except Exception as exc:
            logger.exception(exc)
            error_message = str(exc) or exc.__class__.__name__
            cls._log_chain_step(
                db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="archive-process",
                status="failed",
                reason=error_message,
            )
            cls._exception_record(
                db,
                record_id,
                status_desc="日志拉取流程执行异常",
                error_message=error_message,
                exception_detail=traceback.format_exc(),
            )
            cls._notify_log_pull_record(
                db,
                record,
                status="failed",
                message="日志拉取流程执行异常",
                detail=f"record_id={record.id}, error={error_message}",
            )
        finally:
            if temp_file_path and temp_file_path.exists():
                try:
                    temp_file_path.unlink()
                except Exception:
                    logger.warning(f"删除临时日志压缩包失败: {temp_file_path}")

    @classmethod
    def _trigger_auto_ai_analysis(cls, db: Session, record_id: int) -> None:
        """
        根据日志拉取记录中的自动化配置触发 AI 分析。
        :param db: 数据库会话
        :param record_id: 日志拉取记录ID
        :return: 无
        """
        record = TicketLogPullDao.get_record_by_id(db, record_id)
        if not record or not isinstance(record.command_content, dict):
            cls._log_chain_step(
                db,
                ticket_id=getattr(record, "ticket_id", None),
                record_id=record_id,
                step="auto-ai",
                status="skipped",
                reason="缺少可用的命令内容",
            )
            return
        if not record.ticket_id:
            cls._log_chain_step(
                db,
                ticket_id=None,
                record_id=record_id,
                step="auto-ai",
                status="skipped",
                reason="未关联工单",
            )
            return
        ticket = TicketDao.get_ticket_by_id(db, record.ticket_id)
        if not ticket:
            logger.warning("日志拉取记录[%s] 自动AI触发失败，工单不存在", record_id)
            cls._log_chain_step(
                db,
                ticket_id=record.ticket_id,
                record_id=record_id,
                step="auto-ai",
                status="skipped",
                reason="工单不存在",
            )
            return
        version_key = cls._ensure_ticket_version_key_from_log(db, ticket.ticket_id, record_id)
        record_notify_config = cls._extract_record_notify_config(record)
        automation = record.command_content.get("_automation")
        if isinstance(automation, dict):
            auto_ai_enabled = bool(automation.get("autoAiEnabled"))
            agent_code = str(automation.get("aiAgentCode") or "").strip()
            provider_code = str(automation.get("aiProviderCode") or "").strip()
        else:
            auto_ai_enabled = bool(record.command_content.get("autoAiEnabled"))
            agent_code = str(record.command_content.get("aiAgentCode") or "").strip()
            provider_code = str(record.command_content.get("aiProviderCode") or "").strip()
        if not auto_ai_enabled:
            cls._log_chain_step(
                db,
                ticket_id=record.ticket_id,
                record_id=record_id,
                step="auto-ai",
                status="skipped",
                reason="未启用自动AI",
            )
            return
        if not agent_code and not provider_code:
            logger.warning("日志拉取记录[%s] 已配置自动AI但未填写Provider或Agent", record_id)
            cls._log_chain_step(
                db,
                ticket_id=record.ticket_id,
                record_id=record_id,
                step="auto-ai",
                status="skipped",
                reason="未填写Provider或Agent",
            )
            cls._notify_automation(
                db,
                ticket.ticket_id,
                status="failed",
                message="日志拉取后自动AI已启用，但未配置Provider或Agent，已跳过分析",
                detail=f"record_id={record_id}",
                notify_config=record_notify_config,
            )
            return
        if not version_key:
            logger.warning("日志拉取记录[%s] 自动AI触发失败，工单缺少版本号", record_id)
            cls._log_chain_step(
                db,
                ticket_id=record.ticket_id,
                record_id=record_id,
                step="auto-ai",
                status="skipped",
                reason="工单缺少版本号",
            )
            cls._notify_automation(
                db,
                ticket.ticket_id,
                status="failed",
                message="日志拉取成功但未从日志中提取到版本号，后续AI分析已跳过",
                detail=f"record_id={record_id}",
                notify_config=record_notify_config,
            )
            return
        try:
            from modules.ticket.entity.vo.ticket_vo import TicketAiAnalysisRequestModel
            from modules.ticket.service.ai.ticket_ai_analysis_service import TicketAiAnalysisService

            request = TicketAiAnalysisRequestModel(
                version_key=version_key,
                log_pull_record_id=record.id,
                agent_code=agent_code,
                ai_provider_code=provider_code,
            )
            logger.info(
                "日志拉取记录[%s] 触发自动AI分析 | ticket_id=%s, version_key=%s, agent_code=%s, provider_code=%s",
                record_id,
                record.ticket_id,
                version_key,
                agent_code,
                provider_code,
            )
            result = TicketAiAnalysisService.create_analysis_task_services(db, record.ticket_id, request, None)
            if not result.is_success:
                logger.warning(
                    "日志拉取记录[%s] 自动AI分析未成功提交 | ticket_id=%s, message=%s",
                    record_id,
                    record.ticket_id,
                    result.message,
                )
                cls._notify_automation(
                    db,
                    record.ticket_id,
                    status="failed",
                    message=f"日志拉取后自动AI提交失败：{result.message}",
                    detail=f"record_id={record_id}, version_key={version_key}",
                    notify_config=record_notify_config,
                )
            else:
                cls._log_chain_step(
                    db,
                    ticket_id=record.ticket_id,
                    record_id=record_id,
                    step="auto-ai",
                    status="submitted",
                    reason="自动AI分析已提交",
                    detail={
                        "versionKey": version_key,
                        "agentCode": agent_code,
                        "providerCode": provider_code,
                        "taskId": getattr(result.result, "task_id", None),
                    },
                )
        except Exception as exc:
            logger.exception("日志拉取记录[%s] 触发自动AI分析失败: %s", record_id, exc)
            cls._log_chain_step(
                db,
                ticket_id=record.ticket_id,
                record_id=record_id,
                step="auto-ai",
                status="failed",
                reason=str(exc),
            )
            cls._notify_automation(
                db,
                record.ticket_id,
                status="failed",
                message="日志拉取后自动AI分析触发异常",
                detail=f"record_id={record_id}, error={exc}",
                notify_config=record_notify_config,
            )

    @classmethod
    def _submit_external_request(cls, db: Session, record: TicketLogPullRecord) -> None:
        """
        向外部平台提交日志拉取申请。
        :param db: 数据库会话
        :param record: 日志拉取记录
        :return: 无
        """
        external_config = cls._get_external_config_dict(db)
        request_url = str(external_config.get("insertUrl") or "").strip()
        page_url = str(external_config.get("pageUrl") or "").strip()
        loggable_config = dict(external_config)
        loggable_headers = dict(loggable_config.get("headers") or {})
        if loggable_headers.get("cookie"):
            loggable_headers["cookie"] = "***"
        if loggable_headers.get("authorization"):
            loggable_headers["authorization"] = "***"
        loggable_config["headers"] = loggable_headers
        logger.info(
            "日志拉取外部提交配置: record_id=%s ticket_id=%s config=%s request_url=%s page_url=%s",
            record.id,
            record.ticket_id,
            loggable_config,
            request_url,
            page_url,
        )
        if not request_url:
            cls._log_chain_step(
                db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="submit-external",
                status="skipped",
                reason="外部提交地址未配置",
            )
            raise ValueError("日志拉取外部接口提交地址未配置，请检查 ticket.logPull.external.insertUrl")
        command_content = record.command_content if isinstance(record.command_content, dict) else cls._json_loads(
            record.command_content, {}
        )
        command_content = cls._build_command_content(command_content)
        response = requests.post(
            request_url,
            data={
                "venderId": record.vendor_id,
                "storeId": str(record.store_id or "").strip(),
                "posNo": record.pos_no,
                "commandType": record.command_type,
                "commandDataType": record.command_data_type,
                "commandContent": cls._json_dumps(command_content),
            },
            timeout=(10, 30),
            headers=cls._build_external_request_headers(external_config),
        )
        response.raise_for_status()
        payload = response.json()
        if int(payload.get("code") or 0) != 200 or payload.get("data") is not True:
            cls._log_chain_step(
                db,
                ticket_id=record.ticket_id,
                record_id=record.id,
                step="submit-external",
                status="failed",
                reason=f"外部平台拒绝提交: {payload.get('msg') or payload}",
            )
            raise RuntimeError(f"提交日志拉取申请失败: {payload.get('msg') or payload}")
        cls._update_status(
            db,
            record.id,
            status=TicketLogPullStatus.POLLING.value,
            status_desc="已提交申请，等待外部平台生成压缩包",
            is_error=False,
            update_by="system",
        )
        cls._log_chain_step(
            db,
            ticket_id=record.ticket_id,
            record_id=record.id,
            step="submit-external",
            status="success",
            reason="外部平台提交成功",
            detail={"requestUrl": request_url, "pageUrl": page_url},
        )

    @classmethod
    def _poll_external_result(cls, db: Session, record: TicketLogPullRecord) -> dict[str, Any] | None:
        """
        轮询外部平台列表接口，直到获取成功或失败结果。
        :param db: 数据库会话
        :param record: 日志拉取记录
        :return: 匹配到的外部命令数据
        """
        storage_config = cls._get_storage_config_dict(db)
        deadline = datetime.now() + timedelta(seconds=int(storage_config.get("pollTimeoutSec") or 1800))
        interval_seconds = int(storage_config.get("pollIntervalSec") or 20)
        cls._log_chain_step(
            db,
            ticket_id=record.ticket_id,
            record_id=record.id,
            step="poll-external",
            status="running",
            reason="开始轮询外部平台结果",
            detail={
                "timeoutSeconds": int(storage_config.get("pollTimeoutSec") or 1800),
                "intervalSeconds": interval_seconds,
            },
        )
        while datetime.now() < deadline:
            rows = cls._fetch_external_rows(db, record)
            matched_row = cls._match_external_row(record, rows)
            now = datetime.now()
            TicketLogPullDao.update_record(
                db,
                record.id,
                {
                    "status": TicketLogPullStatus.POLLING.value,
                    "status_desc": "已提交申请，轮询外部平台处理中",
                    "last_polled_at": now,
                    "update_by": "system",
                    "update_time": now,
                },
            )
            db.commit()
            if matched_row:
                external_status = int(matched_row.get("commandStatus") or 0)
                update_data = {
                    "external_command_id": matched_row.get("id"),
                    "external_serial_number": matched_row.get("serialNumber"),
                    "external_command_status": external_status,
                    "external_command_status_desc": matched_row.get("commandStatusDesc"),
                    "command_result_url": cls._extract_command_result_url(matched_row.get("commandResult")),
                    "source_created_at": cls._parse_external_datetime(matched_row),
                    "last_polled_at": now,
                    "update_by": "system",
                    "update_time": now,
                }
                TicketLogPullDao.update_record(db, record.id, update_data)
                db.commit()
                if external_status == 1:
                    cls._log_chain_step(
                        db,
                        ticket_id=record.ticket_id,
                        record_id=record.id,
                        step="poll-external",
                        status="matched",
                        reason="外部平台已生成可下载结果",
                        detail={"externalStatus": external_status, "serialNumber": matched_row.get("serialNumber")},
                    )
                    return matched_row
                if external_status == 2:
                    cls._fail_record(
                        db,
                        record.id,
                        status=TicketLogPullStatus.FAILED.value,
                        status_desc="外部平台日志拉取失败",
                        error_message=str(matched_row.get("errorMsg") or "外部平台返回失败"),
                    )
                    cls._add_ticket_event(
                        db,
                        ticket_id=record.ticket_id,
                        operator_id=None,
                        operator_name="system",
                        content="外部平台日志拉取失败",
                        event_data={"record_id": record.id, "error_message": matched_row.get("errorMsg")},
                    )
                    db.commit()
                    cls._log_chain_step(
                        db,
                        ticket_id=record.ticket_id,
                        record_id=record.id,
                        step="poll-external",
                        status="failed",
                        reason="外部平台返回失败",
                        detail={"errorMessage": matched_row.get("errorMsg")},
                    )
                    return None
            threading.Event().wait(interval_seconds)
        cls._fail_record(
            db,
            record.id,
            status=TicketLogPullStatus.FAILED.value,
            status_desc="轮询外部平台超时",
            error_message="在配置的超时时间内未获取到日志拉取结果",
        )
        cls._log_chain_step(
            db,
            ticket_id=record.ticket_id,
            record_id=record.id,
            step="poll-external",
            status="timeout",
            reason="在配置超时时间内未获取到结果",
        )
        return None

    @classmethod
    def _fetch_external_rows(cls, db: Session, record: TicketLogPullRecord) -> list[dict[str, Any]]:
        """
        查询外部平台日志拉取列表。
        :param db: 数据库会话
        :param record: 日志拉取记录
        :return: 外部平台列表数据
        """
        external_config = cls._get_external_config_dict(db)
        response = requests.get(
            external_config["pageUrl"],
            params={
                "currentPage": 1,
                "pageSize": 20,
                "venderId": record.vendor_id,
                "storeId": str(record.store_id or "").strip(),
                "posNo": record.pos_no,
                "commandType": "",
                "commandStatus": "",
                "_": int(datetime.now().timestamp() * 1000),
            },
            timeout=(10, 30),
            headers=cls._build_external_request_headers(external_config),
        )
        response.raise_for_status()
        payload = response.json()
        if int(payload.get("code") or 0) != 200:
            raise RuntimeError(f"查询日志拉取列表失败: {payload.get('msg') or payload}")
        rows = payload.get("data") or []
        return rows if isinstance(rows, list) else []

    @classmethod
    def _match_external_row(cls, record: TicketLogPullRecord, rows: list[dict[str, Any]]) -> dict[str, Any] | None:
        """
        从外部平台列表结果中匹配当前日志拉取记录。
        :param record: 日志拉取记录
        :param rows: 外部接口返回列表
        :return: 匹配到的记录
        """
        # 外部平台不会回传内部自动化字段，匹配时需要忽略 `_automation`。
        record_command_content = cls._strip_internal_command_content(record.command_content)
        for row in rows:
            if record.external_command_id and int(row.get("id") or 0) == int(record.external_command_id):
                return row
            if int(row.get("venderId") or 0) != int(record.vendor_id):
                continue
            if str(row.get("storeId") or "").strip() != str(record.store_id or "").strip():
                continue
            if int(row.get("posNo") or 0) != int(record.pos_no):
                continue
            if int(row.get("commandType") or 0) != int(record.command_type):
                continue
            if int(row.get("commandDataType") or 0) != int(record.command_data_type):
                continue
            command_content = cls._json_loads(row.get("commandContent"), {})
            if not cls._match_command_content_for_download(record_command_content, command_content):
                continue
            source_created_at = cls._parse_external_datetime(row)
            if source_created_at and source_created_at < record.create_time - timedelta(minutes=5):
                continue
            return row
        return None

    @classmethod
    def _download_archive(cls, record: TicketLogPullRecord, db: Session) -> tuple[Path, int]:
        """
        下载外部压缩包到本地临时文件。
        :param record: 日志拉取记录
        :param db: 数据库会话
        :return: 临时文件路径和文件大小
        """
        config = cls._get_storage_config_dict(db)
        external_config = cls._get_external_config_dict(db)
        timeout_seconds = int(config.get("downloadTimeoutSec") or 300)
        cls._log_chain_step(
            db,
            ticket_id=record.ticket_id,
            record_id=record.id,
            step="download-archive",
            status="running",
            reason="开始下载归档文件",
            detail={"timeoutSeconds": timeout_seconds, "storageMode": record.storage_mode},
        )
        cls.DEFAULT_TEMP_DIR.mkdir(parents=True, exist_ok=True)
        parsed_url = urlparse(str(record.command_result_url or ""))
        source_name = Path(parsed_url.path).name or f"{record.id}.zip"
        temp_fd, temp_name = tempfile.mkstemp(
            prefix=f"{record.id}_",
            suffix=f"_{source_name}",
            dir=cls.DEFAULT_TEMP_DIR,
        )
        os.close(temp_fd)
        temp_file = Path(temp_name)
        total_size = 0
        with requests.get(
            str(record.command_result_url),
            stream=True,
            timeout=(10, timeout_seconds),
            headers=cls._build_external_request_headers(external_config),
        ) as response:
            response.raise_for_status()
            with temp_file.open("wb") as file_obj:
                for chunk in response.iter_content(chunk_size=1024 * 1024):
                    if not chunk:
                        continue
                    file_obj.write(chunk)
                    total_size += len(chunk)
        return temp_file, total_size

    @classmethod
    def _store_archive(cls, record: TicketLogPullRecord, temp_file_path: Path, db: Session) -> str:
        """
        将下载好的压缩包归档到本地或 FTP。
        :param record: 日志拉取记录
        :param temp_file_path: 临时文件路径
        :param db: 数据库会话
        :return: 最终归档地址
        """
        config = cls._get_storage_config_dict(db)
        storage_mode = str(record.storage_mode or config.get("mode") or "local").strip().lower()
        ticket = TicketDao.get_ticket_by_id(db, record.ticket_id)
        ticket_no = getattr(ticket, "ticket_no", None) or str(record.ticket_id)
        relative_path = PurePosixPath(
            "ticket-log-pulls",
            ticket_no,
            datetime.now().strftime("%Y%m%d"),
            temp_file_path.name,
        )
        if storage_mode == "ftp":
            target_path = cls._build_ftp_storage_path(config, relative_path)
            cls._upload_file_to_ftp(config, target_path, temp_file_path)
            return target_path
        local_root = Path(config.get("effectiveLocalDirectory") or cls.DEFAULT_LOCAL_DIR)
        target_path = local_root / Path(relative_path.as_posix())
        target_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(temp_file_path, target_path)
        return str(target_path)

    @classmethod
    def _restore_archive_storage(cls, record: TicketLogPullRecord, temp_file_path: Path, db: Session) -> str:
        """
        将重新下载的压缩包恢复到原归档位置。
        :param record: 日志拉取记录
        :param temp_file_path: 临时文件路径
        :param db: 数据库会话
        :return: 归档位置
        """
        storage_mode = str(
            record.storage_mode or cls._get_storage_config_dict(db).get("mode") or "local"
        ).strip().lower()
        storage_path = str(record.storage_path or "").strip()
        if storage_mode == "ftp" and storage_path:
            config = cls._get_storage_config_dict(db)
            cls._upload_file_to_ftp(config, storage_path, temp_file_path)
            return storage_path
        if storage_mode == "local" and storage_path:
            target_path = Path(storage_path)
            target_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(temp_file_path, target_path)
            return str(target_path)
        return cls._store_archive(record, temp_file_path, db)

    @classmethod
    def _extract_archive_content(
        cls,
        record: TicketLogPullRecord,
        archive_path: Path,
        db: Session,
        begin_time: datetime | None = None,
        end_time: datetime | None = None,
    ) -> dict[str, Any]:
        """
        从压缩包中提取日志文本并压缩入库。
        :param record: 日志拉取记录
        :param archive_path: 压缩包路径
        :param db: 数据库会话
        :param begin_time: 可选的查看开始时间
        :param end_time: 可选的查看结束时间
        :return: 提取结果
        """
        if int(record.command_data_type or TicketLogDataType.LOG.value) == TicketLogDataType.DB.value:
            return {
                "archive_entry_count": cls._count_archive_entries(archive_path),
                "matched_entry_count": 0,
                "content_char_count": 0,
                "content_truncated": False,
                "compressed_content": None,
                "content_summary": "DB 拉取完成，当前版本不解析数据库文件文本内容。",
            }

        config = cls._get_storage_config_dict(db)
        max_content_chars = int(config.get("maxContentChars") or 500000)
        timestamp_start = begin_time if begin_time is not None else cls._parse_datetime(record.log_begin_time)
        timestamp_end = end_time if end_time is not None else cls._parse_datetime(record.log_end_time)
        archive_entry_count = 0
        matched_entry_count = 0
        content_parts: list[str] = []
        current_char_count = 0

        with zipfile.ZipFile(archive_path) as archive:
            entry_names = [name for name in archive.namelist() if not name.endswith("/")]
            archive_entry_count = len(entry_names)
            target_entry_names = [name for name in entry_names if cls._is_target_log_entry(name)]
            if not target_entry_names:
                return {
                    "archive_entry_count": archive_entry_count,
                    "matched_entry_count": 0,
                    "content_char_count": 0,
                    "content_truncated": False,
                    "compressed_content": None,
                    "content_summary": f"压缩包共 {archive_entry_count} 个文件，未发现 *_pos.log* 日志文件。",
                }
            for entry_name in cls._sort_archive_entry_names(target_entry_names):
                matched_entry_count, current_char_count = cls._extract_entry_logs(
                    archive=archive,
                    entry_name=entry_name,
                    record=record,
                    timestamp_start=timestamp_start,
                    timestamp_end=timestamp_end,
                    content_parts=content_parts,
                    current_char_count=current_char_count,
                    current_matched_count=matched_entry_count,
                    max_content_chars=max_content_chars,
                )

        full_text = "\n\n".join(part for part in content_parts if part)
        return {
            "archive_entry_count": archive_entry_count,
            "matched_entry_count": matched_entry_count,
            "content_char_count": len(full_text),
            "content_truncated": False,
            "compressed_content": cls._compress_text(full_text) if full_text else None,
            "content_summary": cls._build_content_summary(
                archive_entry_count=archive_entry_count,
                matched_entry_count=matched_entry_count,
                content_char_count=len(full_text),
                content_truncated=False,
            ),
        }

    @classmethod
    def _sort_archive_entry_names(cls, entry_names: list[str]) -> list[str]:
        """
        按压缩包日志文件尾号排序，尾号越小越接近当前时间，因此需要按尾号逆序拼接。
        :param entry_names: 压缩包中的日志文件名列表
        :return: 排序后的文件名列表
        """
        def _entry_sort_key(entry_name: str) -> tuple[int, int, str]:
            file_name = PurePosixPath(entry_name).name
            match = re.search(r"(\d+)(?=\D*$)", file_name)
            suffix_index = int(match.group(1)) if match else -1
            return (1 if match else 0, suffix_index, file_name.lower())

        return sorted(entry_names, key=_entry_sort_key, reverse=True)

    @classmethod
    def _match_command_content_for_download(
        cls, local_command_content: dict[str, Any] | None, remote_command_content: dict[str, Any] | None
    ) -> bool:
        """
        比较本地与远端 commandContent 中用于下载命中的参数。
        :param local_command_content: 本地记录中的命令内容
        :param remote_command_content: 外部平台返回的命令内容
        :return: 是否匹配
        """
        local_content = cls._strip_internal_command_content(local_command_content)
        remote_content = cls._strip_internal_command_content(remote_command_content)
        local_fields = {
            key: str(local_content.get(key)).strip()
            for key in ("modifyTime", "path")
            if str(local_content.get(key) or "").strip()
        }
        remote_fields = {
            key: str(remote_content.get(key)).strip()
            for key in ("modifyTime", "path")
            if str(remote_content.get(key) or "").strip()
        }
        if len(local_fields) != len(remote_fields):
            return False
        return local_fields == remote_fields

    @classmethod
    def _extract_entry_logs(
        cls,
        *,
        archive: zipfile.ZipFile,
        entry_name: str,
        record: TicketLogPullRecord,
        timestamp_start: datetime | None,
        timestamp_end: datetime | None,
        content_parts: list[str],
        current_char_count: int,
        current_matched_count: int,
        max_content_chars: int,
    ) -> tuple[int, int]:
        """
        解析压缩包中的单个日志文件。
        :param archive: 压缩包对象
        :param entry_name: 文件名
        :param record: 日志拉取记录
        :param timestamp_start: 开始时间
        :param timestamp_end: 结束时间
        :param content_parts: 已匹配的文本片段列表
        :param current_char_count: 当前已累积的字符数
        :param current_matched_count: 当前已命中的日志条目数
        :param max_content_chars: 最大保留字符数
        :return: 更新后的命中数和字符数
        """
        encoding = cls._detect_archive_entry_encoding(archive, entry_name)
        with archive.open(entry_name, "r") as binary_file:
            text_file = io.TextIOWrapper(binary_file, encoding=encoding, errors="replace")
            current_lines: list[str] = []
            current_timestamp: datetime | None = None
            matched_count = current_matched_count
            char_count = current_char_count
            for raw_line in text_file:
                line = raw_line.rstrip("\r\n")
                match = cls.TIMESTAMP_PATTERN.match(line)
                if match:
                    matched_count, char_count = cls._flush_log_entry(
                        current_lines=current_lines,
                        current_timestamp=current_timestamp,
                        begin_time=timestamp_start,
                        end_time=timestamp_end,
                        matched_count=matched_count,
                        content_parts=content_parts,
                        char_count=char_count,
                        max_content_chars=max_content_chars,
                    )
                    current_timestamp = datetime.strptime(match.group(1), "%Y-%m-%d %H:%M:%S,%f")
                    file_name = os.path.basename(entry_name)
                    current_lines = [f"[{file_name}]{line}"]
                    continue
                if current_lines:
                    current_lines.append(line)
            matched_count, char_count = cls._flush_log_entry(
                current_lines=current_lines,
                current_timestamp=current_timestamp,
                begin_time=timestamp_start,
                end_time=timestamp_end,
                matched_count=matched_count,
                content_parts=content_parts,
                char_count=char_count,
                max_content_chars=max_content_chars,
            )
            return matched_count, char_count

    @classmethod
    def _resolve_view_log_time_range(
        cls, query: TicketLogPullContentQueryModel | None
    ) -> tuple[datetime | None, datetime | None]:
        """
        将查看日志参数统一解析为开始和结束时间。
        :param query: 查看参数
        :return: 开始时间和结束时间
        """
        if not query:
            return None, None
        log_begin_time = cls._parse_datetime(query.log_begin_time)
        log_end_time = cls._parse_datetime(query.log_end_time)
        if log_begin_time or log_end_time:
            return log_begin_time, log_end_time

        point_time = cls._parse_datetime(query.log_point_time)
        if not point_time:
            return None, None
        before_minutes = max(int(query.range_before_minutes or 0), 0)
        after_minutes = max(int(query.range_after_minutes or 0), 0)
        return (
            point_time - timedelta(minutes=before_minutes),
            point_time + timedelta(minutes=after_minutes),
        )

    @classmethod
    def _resolve_archive_source_for_view(cls, record: TicketLogPullRecord, db: Session) -> tuple[Path | None, bool]:
        """
        解析查看日志时可用的压缩包来源。
        :param record: 日志拉取记录
        :param db: 数据库会话
        :return: 压缩包路径和是否需要清理临时文件
        """
        storage_path = str(record.storage_path or "").strip()
        if storage_path:
            local_path = Path(storage_path)
            if local_path.exists() and local_path.is_file():
                return local_path, False
        storage_mode = str(record.storage_mode or "").strip().lower()
        if storage_mode == "ftp" and storage_path:
            try:
                return cls._download_file_from_ftp_to_temp(db, storage_path), True
            except Exception as exc:
                logger.warning("从 FTP 下载日志压缩包失败: %s", exc)
        if str(record.command_result_url or "").strip():
            try:
                temp_file, _ = cls._download_archive(record, db)
                return temp_file, True
            except Exception as exc:
                logger.warning("从外部地址重新下载日志压缩包失败: %s", exc)
        return None, False

    @classmethod
    def _resolve_archive_source_for_download(
        cls, record: TicketLogPullRecord, db: Session, *, source: str = "auto"
    ) -> tuple[Path | None, bool]:
        """
        按下载来源解析压缩包文件。
        :param record: 日志拉取记录
        :param db: 数据库会话
        :param source: 下载来源，auto/service/original
        :return: 压缩包路径和是否需要清理临时文件
        """
        normalized_source = str(source or "auto").strip().lower()
        storage_path = str(record.storage_path or "").strip()
        storage_mode = str(record.storage_mode or "").strip().lower()
        if normalized_source in {"auto", "service"} and storage_path:
            local_path = Path(storage_path)
            if local_path.exists() and local_path.is_file():
                return local_path, False
            if storage_mode == "ftp":
                try:
                    return cls._download_file_from_ftp_to_temp(db, storage_path), True
                except Exception as exc:
                    logger.warning("从 FTP 下载日志压缩包失败: %s", exc)
            if normalized_source == "service":
                return None, False
        if normalized_source in {"auto", "original"} and str(record.command_result_url or "").strip():
            try:
                temp_file, _ = cls._download_archive(record, db)
                return temp_file, True
            except Exception as exc:
                logger.warning("从外部地址重新下载日志压缩包失败: %s", exc)
        return None, False

    @classmethod
    def _build_download_file_name(cls, record: TicketLogPullRecord, archive_path: Path | None = None) -> str:
        """
        生成日志压缩包下载文件名。
        :param record: 日志拉取记录
        :param archive_path: 可选的归档文件路径
        :return: 下载文件名
        """
        candidates = [
            str(record.download_file_name or "").strip(),
            Path(str(record.storage_path or "")).name if str(record.storage_path or "").strip() else "",
            Path(urlparse(str(record.command_result_url or "")).path).name
            if str(record.command_result_url or "").strip()
            else "",
            archive_path.name if archive_path else "",
            f"ticket_log_pull_{record.id}.zip",
        ]
        for candidate in candidates:
            if not candidate:
                continue
            resolved = Path(candidate).name
            if resolved:
                return resolved
        return f"ticket_log_pull_{record.id}.zip"

    @classmethod
    def _download_file_from_ftp_to_temp(cls, db: Session, remote_path: str) -> Path:
        """
        从 FTP 下载文件到临时路径。
        :param db: 数据库会话
        :param remote_path: FTP 文件路径
        :return: 临时文件路径
        """
        config = cls._get_storage_config_dict(db)
        cls.DEFAULT_TEMP_DIR.mkdir(parents=True, exist_ok=True)
        suffix = Path(PurePosixPath(remote_path).name).suffix or ".zip"
        temp_fd, temp_name = tempfile.mkstemp(prefix="ticket-log-view-", suffix=suffix, dir=cls.DEFAULT_TEMP_DIR)
        os.close(temp_fd)
        temp_path = Path(temp_name)
        ftp = cls._connect_ftp(config)
        try:
            with temp_path.open("wb") as file_obj:
                ftp.retrbinary(f"RETR {remote_path}", file_obj.write)
            return temp_path
        except ftp_errors as exc:
            try:
                temp_path.unlink(missing_ok=True)
            except Exception:
                pass
            raise RuntimeError(f"FTP 下载失败: {exc}") from exc
        finally:
            try:
                ftp.quit()
            except Exception:
                try:
                    ftp.close()
                except Exception:
                    pass

    @classmethod
    def _delete_record_storage(cls, record: TicketLogPullRecord, db: Session) -> None:
        """
        删除日志拉取记录对应的文件数据。
        :param record: 日志拉取记录
        :param db: 数据库会话
        :return: 无
        """
        storage_path = str(record.storage_path or "").strip()
        if not storage_path:
            return
        storage_mode = (
            str(record.storage_mode or cls._get_storage_config_dict(db).get("mode") or "local").strip().lower()
        )
        if storage_mode == "ftp":
            ftp = cls._connect_ftp(cls._get_storage_config_dict(db))
            try:
                ftp.delete(storage_path)
            except error_perm as exc:
                message = str(exc)
                if "550" in message or "not found" in message.lower() or "no such file" in message.lower():
                    return
                raise RuntimeError(f"FTP 删除失败: {exc}") from exc
            except ftp_errors as exc:
                raise RuntimeError(f"FTP 删除失败: {exc}") from exc
            finally:
                try:
                    ftp.quit()
                except Exception:
                    try:
                        ftp.close()
                    except Exception:
                        pass
            return
        local_path = Path(storage_path)
        if local_path.exists():
            try:
                local_path.unlink()
            except Exception as exc:
                raise RuntimeError(f"删除本地文件失败: {exc}") from exc

    @staticmethod
    def _build_view_fallback_summary(base_summary: str | None, reason: str) -> str:
        """
        构建查看日志时的降级摘要。
        :param base_summary: 原始摘要
        :param reason: 降级原因
        :return: 合成摘要
        """
        base_text = str(base_summary or "").strip()
        reason_text = str(reason or "").strip()
        if base_text and reason_text:
            return f"{base_text}（{reason_text}）"
        return base_text or reason_text

    @classmethod
    def _flush_log_entry(
        cls,
        *,
        current_lines: list[str],
        current_timestamp: datetime | None,
        begin_time: datetime | None,
        end_time: datetime | None,
        matched_count: int,
        content_parts: list[str],
        char_count: int,
        max_content_chars: int,
    ) -> tuple[int, int]:
        """
        将当前正在累计的日志条目按时间范围写入结果集。
        :param current_lines: 当前条目文本行
        :param current_timestamp: 当前条目时间
        :param begin_time: 开始时间
        :param end_time: 结束时间
        :param matched_count: 当前命中数
        :param content_parts: 结果集
        :param char_count: 当前字符数
        :param max_content_chars: 最大字符数
        :return: 更新后的命中数和字符数
        """
        if not current_lines or current_timestamp is None:
            return matched_count, char_count
        if begin_time and current_timestamp < begin_time:
            return matched_count, char_count
        if end_time and current_timestamp > end_time:
            return matched_count, char_count
        entry_text = "\n".join(current_lines)
        if not entry_text.strip():
            return matched_count, char_count
        separator_length = 2 if content_parts else 0
        projected_length = char_count + len(entry_text) + separator_length
        if False and projected_length > max_content_chars:  # 这里暂时取消长度判断
            raise TicketLogContentTooLargeError(
                f"日志内容超过入库上限 {max_content_chars} 字符，请缩小时间范围后重新提交"
            )
        content_parts.append(entry_text)
        return matched_count + 1, projected_length

    @classmethod
    def _detect_archive_entry_encoding(cls, archive: zipfile.ZipFile, entry_name: str) -> str:
        """
        尝试识别压缩包中文本文件编码。
        :param archive: 压缩包对象
        :param entry_name: 文件名
        :return: 编码名称
        """
        with archive.open(entry_name, "r") as binary_file:
            sample = binary_file.read(4096)
        for encoding in ("utf-8", "gb18030", "latin-1"):
            try:
                sample.decode(encoding)
                return encoding
            except Exception:
                continue
        return "utf-8"

    @classmethod
    def _count_archive_entries(cls, archive_path: Path) -> int:
        """
        统计压缩包内文件数量。
        :param archive_path: 压缩包路径
        :return: 文件数量
        """
        with zipfile.ZipFile(archive_path) as archive:
            return len([name for name in archive.namelist() if not name.endswith("/")])

    @classmethod
    def _build_content_summary(
        cls, *, archive_entry_count: int, matched_entry_count: int, content_char_count: int, content_truncated: bool
    ) -> str:
        """
        生成日志内容摘要。
        :param archive_entry_count: 压缩包文件数
        :param matched_entry_count: 命中条目数
        :param content_char_count: 文本字符数
        :param content_truncated: 是否截断
        :return: 摘要文本
        """
        summary = (
            f"压缩包共 {archive_entry_count} 个文件，命中 {matched_entry_count} 条日志，"
            f"文本长度 {content_char_count} 字符"
        )
        if content_truncated:
            summary += "，已按配置截断"
        return summary

    @staticmethod
    def _has_explicit_log_time_range(payload: TicketLogPullCreateModel) -> bool:
        """
        判断提交参数是否包含明确的日志截取时间范围。
        :param payload: 日志拉取参数
        :return: 是否配置了时间范围
        """
        fields_set = getattr(payload, "model_fields_set", set()) or set()
        has_direct_range = any(field in fields_set for field in ("log_begin_time", "log_end_time"))
        has_point_range = any(
            field in fields_set
            for field in ("log_point_time", "range_before_minutes", "range_after_minutes")
        )
        return has_direct_range or has_point_range

    @classmethod
    def _compress_text(cls, text: str) -> str | None:
        """
        将文本按 gzip+base64 压缩。
        :param text: 原始文本
        :return: 压缩后的字符串
        """
        if not text:
            return None
        compressed = gzip.compress(text.encode("utf-8"))
        return base64.b64encode(compressed).decode("ascii")

    @classmethod
    def _decompress_text(cls, encoded_text: str | None) -> str:
        """
        将 gzip+base64 文本解压为普通字符串。
        :param encoded_text: 压缩后的字符串
        :return: 解压文本
        """
        if not encoded_text:
            return ""
        return gzip.decompress(base64.b64decode(encoded_text.encode("ascii"))).decode("utf-8")

    @classmethod
    def _build_command_content(cls, payload: TicketLogPullCreateModel | dict[str, Any]) -> dict[str, Any]:
        """
        根据页面输入构造外部接口 commandContent。
        :param payload: 页面请求参数
        :return: commandContent 字典
        """
        source = payload if isinstance(payload, dict) else payload.model_dump(by_alias=True)
        command_content: dict[str, Any] = {
            "fileMaxSize": str(source.get("fileMaxSize") or source.get("file_max_size") or 500),
            "zipMaxSize": str(source.get("zipMaxSize") or source.get("zip_max_size") or 500),
        }
        modify_time = source.get("modifyTime") or source.get("modify_time")
        if modify_time:
            command_content["modifyTime"] = str(modify_time)[:10]
        if str(source.get("path") or "").strip():
            command_content["path"] = str(source.get("path")).strip()
        command_content["commandDataType"] = int(
            source.get("commandDataType") or source.get("command_data_type") or 1
        )
        command_content["storageMode"] = (
            str(source.get("storageMode") or source.get("storage_mode") or "").strip() or None
        )
        point_time = cls._parse_datetime(source.get("logPointTime") or source.get("log_point_time"))
        begin_time = cls._parse_datetime(source.get("logBeginTime") or source.get("log_begin_time"))
        end_time = cls._parse_datetime(source.get("logEndTime") or source.get("log_end_time"))
        if point_time:
            command_content["timeRangeMode"] = "point"
            command_content["logPointTime"] = point_time.isoformat(sep=" ")
            command_content["rangeBeforeMinutes"] = int(
                source.get("rangeBeforeMinutes") or source.get("range_before_minutes") or 0
            )
            command_content["rangeAfterMinutes"] = int(
                source.get("rangeAfterMinutes") or source.get("range_after_minutes") or 0
            )
        elif begin_time or end_time:
            command_content["timeRangeMode"] = "between"
            if begin_time:
                command_content["logBeginTime"] = begin_time.isoformat(sep=" ")
            if end_time:
                command_content["logEndTime"] = end_time.isoformat(sep=" ")
        return command_content

    @staticmethod
    def _strip_internal_command_content(command_content: dict[str, Any] | None) -> dict[str, Any]:
        """
        移除 commandContent 中仅供系统内部使用的扩展字段。
        :param command_content: 原始 commandContent
        :return: 发送给外部接口的 commandContent
        """
        content = dict(command_content or {})
        content.pop("_automation", None)
        return content

    @classmethod
    def _resolve_log_time_range(
        cls, payload: TicketLogPullCreateModel
    ) -> tuple[datetime | None, datetime | None]:
        """
        将日志时间范围统一解析为开始和结束时间。
        :param payload: 页面请求参数
        :return: 开始时间和结束时间
        """
        log_begin_time = cls._parse_datetime(payload.log_begin_time)
        log_end_time = cls._parse_datetime(payload.log_end_time)
        if log_begin_time and log_end_time:
            return log_begin_time, log_end_time

        point_time = cls._parse_datetime(payload.log_point_time)
        if not point_time:
            return None, None

        before_minutes = max(int(payload.range_before_minutes or 0), 0)
        after_minutes = max(int(payload.range_after_minutes or 0), 0)
        return (
            point_time - timedelta(minutes=before_minutes),
            point_time + timedelta(minutes=after_minutes),
        )

    @classmethod
    def _build_retry_payload(cls, record: TicketLogPullRecord) -> TicketLogPullCreateModel | None:
        """
        从历史记录恢复重新拉取所需的提交参数。
        :param record: 日志拉取记录
        :return: 可重新提交的创建模型
        """
        command_content = (
            dict(record.command_content)
            if isinstance(record.command_content, dict)
            else cls._json_loads(record.command_content, {})
        )
        if not isinstance(command_content, dict):
            command_content = {}
        time_range_mode = str(
            cls._first_present_value(command_content, "timeRangeMode", "time_range_mode") or ""
        ).strip().lower()
        payload_data: dict[str, Any] = {
            "vendorId": record.vendor_id,
            "storeId": str(record.store_id or "").strip(),
            "posNo": record.pos_no,
            "commandDataType": record.command_data_type,
            "fileMaxSize": cls._parse_positive_int(
                cls._first_present_value(command_content, "fileMaxSize", "file_max_size"), 500
            ),
            "zipMaxSize": cls._parse_positive_int(
                cls._first_present_value(command_content, "zipMaxSize", "zip_max_size"), 500
            ),
            "storageMode": record.storage_mode,
        }
        modify_time = cls._first_present_value(command_content, "modifyTime", "modify_time")
        path = cls._first_present_value(command_content, "path")
        if modify_time:
            payload_data["modifyTime"] = str(modify_time)[:10]
        if path:
            payload_data["path"] = str(path).strip()

        point_time_value = cls._first_present_value(command_content, "logPointTime", "log_point_time")
        if time_range_mode == "point" and point_time_value not in (None, ""):
            before_minutes = cls._first_present_value(
                command_content, "rangeBeforeMinutes", "range_before_minutes"
            )
            after_minutes = cls._first_present_value(command_content, "rangeAfterMinutes", "range_after_minutes")
            payload_data.update(
                {
                    "logPointTime": point_time_value,
                    "rangeBeforeMinutes": before_minutes,
                    "rangeAfterMinutes": after_minutes,
                }
            )
        else:
            # 历史记录可能没有配置日志截取范围，不能把 None 显式传给模型，否则会被判定为范围缺失。
            begin_time = (
                cls._first_present_value(command_content, "logBeginTime", "log_begin_time")
                or record.log_begin_time
            )
            end_time = (
                cls._first_present_value(command_content, "logEndTime", "log_end_time")
                or record.log_end_time
            )
            if begin_time or end_time:
                payload_data["logBeginTime"] = begin_time
                payload_data["logEndTime"] = end_time
        notify_config = command_content.get("notifyConfig") or command_content.get("notify_config")
        if not isinstance(notify_config, dict):
            automation = (
                command_content.get("_automation")
                if isinstance(command_content.get("_automation"), dict)
                else {}
            )
            notify_config = automation.get("notifyConfig") or automation.get("notify_config")
        if isinstance(notify_config, dict):
            payload_data["notifyConfig"] = notify_config
        automation = command_content.get("_automation") if isinstance(command_content.get("_automation"), dict) else {}
        if automation:
            payload_data["autoAiEnabled"] = automation.get("autoAiEnabled")
            payload_data["aiAgentCode"] = automation.get("aiAgentCode")
            payload_data["aiProviderCode"] = automation.get("aiProviderCode")
        try:
            return TicketLogPullCreateModel.model_validate(payload_data)
        except Exception as exc:
            logger.warning("恢复日志拉取参数失败: %s", exc)
            return None

    @classmethod
    def _normalize_command_content(cls, content: Any) -> dict[str, str]:
        """
        标准化 commandContent，避免键顺序差异导致匹配失败。
        :param content: 原始内容
        :return: 标准化后的字典
        """
        raw = dict(content or {}) if isinstance(content, dict) else cls._json_loads(content, {})
        return {str(key): str(value) for key, value in sorted(raw.items(), key=lambda item: item[0])}

    @classmethod
    def _parse_datetime(cls, value: date | datetime | str | None) -> datetime | None:
        """
        解析日期/时间输入。
        :param value: 日期、时间或字符串
        :return: datetime 对象
        """
        if value in (None, ""):
            return None
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, time.min)
        text = str(value).strip()
        if not text:
            return None
        for formatter in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(text[:19] if "T" in text or " " in text else text[:10], formatter)
                if formatter == "%Y-%m-%d":
                    return datetime.combine(parsed.date(), time.min)
                return parsed
            except Exception:
                continue
        return datetime.fromisoformat(text.replace("Z", ""))

    @classmethod
    def _parse_external_datetime(cls, row: dict[str, Any]) -> datetime | None:
        """
        解析外部接口返回的创建时间。
        :param row: 外部接口单条数据
        :return: datetime 对象
        """
        created_value = row.get("created")
        if created_value not in (None, ""):
            try:
                return datetime.fromtimestamp(int(created_value) / 1000)
            except Exception:
                pass
        created_formatted = row.get("createdF")
        if created_formatted:
            try:
                return datetime.strptime(str(created_formatted), "%Y-%m-%d %H:%M:%S")
            except Exception:
                return None
        return None

    @classmethod
    def _is_target_log_entry(cls, entry_name: str) -> bool:
        """
        判断压缩包条目是否属于当前需要解析的 POS 日志文件。
        :param entry_name: 压缩包内文件路径
        :return: 是否为目标日志
        """
        file_name = PurePosixPath(entry_name).name.lower()
        return "_pos.log" in file_name

    @classmethod
    def _extract_command_result_url(cls, command_result: Any) -> str | None:
        """
        提取外部命令结果中的压缩包地址。
        :param command_result: 外部 commandResult 字段
        :return: 压缩包地址
        """
        payload = command_result if isinstance(command_result, dict) else cls._json_loads(command_result, {})
        return str(payload.get("url") or "").strip() or None

    @classmethod
    def _update_status(cls, db: Session, record_id: int, **kwargs) -> None:
        """
        更新日志拉取记录状态字段并提交。
        :param db: 数据库会话
        :param record_id: 记录ID
        :param kwargs: 更新字段
        :return: 无
        """
        kwargs["update_time"] = kwargs.get("update_time") or datetime.now()
        TicketLogPullDao.update_record(db, record_id, kwargs)
        db.commit()

    @classmethod
    def _fail_record(
        cls, db: Session, record_id: int, *, status: str, status_desc: str, error_message: str
    ) -> None:
        """
        将记录标记为失败状态。
        :param db: 数据库会话
        :param record_id: 记录ID
        :param status: 失败状态编码
        :param status_desc: 状态说明
        :param error_message: 错误信息
        :return: 无
        """
        cls._update_status(
            db,
            record_id,
            status=status,
            status_desc=status_desc,
            is_error=True,
            error_message=error_message,
            finished_at=datetime.now(),
            update_by="system",
        )

    @classmethod
    def _exception_record(
        cls,
        db: Session,
        record_id: int,
        *,
        status_desc: str,
        error_message: str,
        exception_detail: str,
    ) -> None:
        """
        将记录标记为程序异常状态。
        :param db: 数据库会话
        :param record_id: 记录ID
        :param status_desc: 状态说明
        :param error_message: 错误信息
        :param exception_detail: 异常堆栈
        :return: 无
        """
        try:
            cls._update_status(
                db,
                record_id,
                status=TicketLogPullStatus.EXCEPTION.value,
                status_desc=status_desc,
                is_error=True,
                error_message=error_message,
                exception_detail=exception_detail,
                finished_at=datetime.now(),
                update_by="system",
            )
            record = TicketLogPullDao.get_record_by_id(db, record_id)
            if record:
                cls._add_ticket_event(
                    db,
                    ticket_id=record.ticket_id,
                    operator_id=None,
                    operator_name="system",
                    content="日志拉取流程异常",
                    event_data={"record_id": record.id, "error_message": error_message},
                )
                db.commit()
        except Exception as exc:
            db.rollback()
            logger.exception(exc)

    @classmethod
    def _get_storage_config_dict(cls, db: Session) -> dict[str, Any]:
        """
        读取日志拉取存储配置字典。
        :param db: 数据库会话
        :return: 配置字典
        """
        cls.ensure_param_config_rows(db)
        config_row = TicketLogPullDao.get_storage_config_row(db)
        payload = cls._json_loads(getattr(config_row, "config_value", None), {})
        return cls._normalize_storage_config(payload)

    @classmethod
    def _get_external_config_dict(cls, db: Session) -> dict[str, Any]:
        """
        读取日志拉取外部接口配置字典。
        :param db: 数据库会话
        :return: 外部接口配置
        """
        cls.ensure_param_config_rows(db)
        config_row = TicketLogPullDao.get_external_config_row(db)
        payload = cls._json_loads(getattr(config_row, "config_value", None), {})
        return cls._normalize_external_config(payload)

    @classmethod
    def _add_ticket_event(
        cls,
        db: Session,
        *,
        ticket_id: int | None,
        operator_id: int | None,
        operator_name: str,
        content: str,
        event_data: dict[str, Any] | None,
    ) -> None:
        """
        追加工单日志分析事件，方便时间线复盘。
        :param db: 数据库会话
        :param ticket_id: 工单ID
        :param operator_id: 操作人ID
        :param operator_name: 操作人名称
        :param content: 事件说明
        :param event_data: 结构化数据
        :return: 无
        """
        if not ticket_id:
            return
        TicketDao.add_event(
            db,
            TicketEvent(
                ticket_id=ticket_id,
                event_type=TicketEventType.LOG_ANALYSIS.value,
                operator_id=operator_id,
                operator_name=operator_name,
                content=content,
                event_data=event_data,
                create_time=datetime.now(),
            ),
        )

    @classmethod
    def _to_record_summary(cls, record: TicketLogPullRecord) -> dict[str, Any]:
        """
        将记录对象转换为最新摘要。
        :param record: 日志拉取记录
        :return: 摘要字典
        """
        command_content = record.command_content if isinstance(record.command_content, dict) else {}
        modify_time = (
            getattr(record, "modify_time", None)
            or command_content.get("modifyTime")
            or command_content.get("modify_time")
        )
        return {
            "id": record.id,
            "status": record.status,
            "statusDesc": record.status_desc,
            "isError": record.is_error,
            "errorMessage": record.error_message,
            "vendorId": record.vendor_id,
            "storeId": str(record.store_id or "").strip() or None,
            "posNo": record.pos_no,
            "modifyTime": modify_time,
            "createTime": record.create_time,
        }

    @classmethod
    def _to_record_list_item(cls, row: dict[str, Any] | TicketLogPullRecord) -> dict[str, Any]:
        """
        将记录对象转换为前端列表项。
        :param row: 数据库对象或已转换字典
        :return: 列表项字典
        """
        payload = row if isinstance(row, dict) else CamelCaseUtil.transform_result(row)
        command_content = payload.get("commandContent")
        store_id = payload.get("storeId")
        if store_id is not None:
            payload["storeId"] = str(store_id).strip() or None
        if not payload.get("modifyTime") and isinstance(command_content, dict):
            payload["modifyTime"] = command_content.get("modifyTime")
        payload["hasContent"] = bool(payload.get("compressedContent"))
        payload.pop("compressedContent", None)
        payload.pop("exceptionDetail", None)
        return TicketLogPullListItemModel.model_validate(payload).model_dump(by_alias=True)

    @classmethod
    def _enrich_record_list_items(cls, db: Session, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        """
        为日志拉取列表补充工单信息。
        :param db: 数据库会话
        :param rows: 日志拉取记录列表
        :return: 补充后的列表
        """
        if not rows:
            return []
        ticket_ids = [item.get("ticketId") for item in rows if item.get("ticketId")]
        ticket_map = {ticket.ticket_id: ticket for ticket in TicketDao.get_tickets_by_ids(db, ticket_ids)}
        enriched_rows: list[dict[str, Any]] = []
        for item in rows:
            payload = dict(item or {})
            ticket_id = payload.get("ticketId")
            ticket = ticket_map.get(ticket_id)
            if ticket:
                payload["ticketNo"] = ticket.ticket_no
                payload["ticketTitle"] = ticket.title
                payload["projectName"] = ticket.merchant_name
                payload["moduleName"] = ticket.module_name
            enriched_rows.append(cls._to_record_list_item(payload))
        return enriched_rows

    @classmethod
    def _build_ftp_storage_path(cls, config: dict[str, Any], relative_path: PurePosixPath) -> str:
        """
        构建 FTP 归档路径。
        :param config: 存储配置
        :param relative_path: 相对路径
        :return: FTP 完整路径
        """
        base_dir = str(config.get("ftp", {}).get("baseDir") or "").strip().replace("\\", "/")
        target_path = PurePosixPath(base_dir or "/") / relative_path
        return target_path.as_posix()

    @classmethod
    def _connect_ftp(cls, config: dict[str, Any]) -> FTP:
        """
        建立 FTP 连接。
        :param config: 存储配置
        :return: FTP 连接对象
        """
        ftp_config = config.get("ftp") if isinstance(config.get("ftp"), dict) else {}
        host = str(ftp_config.get("host") or "").strip()
        username = str(ftp_config.get("username") or "").strip()
        if not host:
            raise RuntimeError("FTP 存储未配置 host")
        if not username:
            raise RuntimeError("FTP 存储未配置 username")
        ftp = FTP()
        ftp.encoding = str(ftp_config.get("encoding") or "utf-8")
        ftp.connect(
            host=host,
            port=int(ftp_config.get("port") or 21),
            timeout=max(int(ftp_config.get("timeoutSec") or 15), 1),
        )
        ftp.login(username, str(ftp_config.get("password") or ""))
        ftp.set_pasv(bool(ftp_config.get("passive", True)))
        return ftp

    @classmethod
    def _ensure_ftp_directory(cls, ftp: FTP, directory: str) -> None:
        """
        递归创建 FTP 目录。
        :param ftp: FTP 连接
        :param directory: 目录路径
        :return: 无
        """
        normalized = PurePosixPath(directory or "/")
        current_path = ""
        for part in normalized.parts:
            if part in ("", ".", "/"):
                continue
            current_path = f"{current_path}/{part}" if current_path else part
            try:
                ftp.mkd(current_path)
            except error_perm as exc:
                message = str(exc)
                if "File exists" in message or message.startswith("550") or message.startswith("521"):
                    continue
                raise

    @classmethod
    def _upload_file_to_ftp(cls, config: dict[str, Any], target_path: str, source_file: Path) -> None:
        """
        将本地文件上传到 FTP。
        :param config: 存储配置
        :param target_path: FTP 目标路径
        :param source_file: 本地源文件
        :return: 无
        """
        ftp = cls._connect_ftp(config)
        try:
            directory = str(PurePosixPath(target_path).parent)
            if directory and directory not in (".", "/"):
                cls._ensure_ftp_directory(ftp, directory)
            with source_file.open("rb") as file_obj:
                ftp.storbinary(f"STOR {target_path}", file_obj)
        except ftp_errors as exc:
            raise RuntimeError(f"FTP 上传失败: {exc}") from exc
        finally:
            try:
                ftp.quit()
            except Exception:
                try:
                    ftp.close()
                except Exception:
                    pass
