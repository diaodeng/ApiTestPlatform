from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from module_admin.entity.vo.user_vo import CurrentUserModel
from module_hrm.entity.do.module_do import HrmModule
from module_hrm.entity.do.project_do import HrmProject
from module_hrm.enums.enums import QtrDataStatusEnum
from modules.ticket.dao.ticket_dao import TicketDao
from modules.ticket.entity.do.ticket_do import Ticket, TicketEvent, TicketRca, TicketStatusHistory
from modules.ticket.enums.ticket_enums import TicketEventType, TicketStatus
from modules.ticket.service.ai.ticket_embedding_service import TicketEmbeddingService
from utils.snowflake import snowIdWorker

IMPORT_HEADERS = [
    "工单号",
    "标题",
    "描述",
    "状态",
    "所属项目",
    "所属模块",
    "问题分类",
    "对方优先级",
    "内部优先级",
    "严重等级",
    "来源",
    "提单人",
    "当前处理人",
    "是否真实问题",
    "根因分类",
    "根因详情",
    "触发原因",
    "影响范围",
    "复现步骤",
    "排查过程",
    "解决方案",
    "验证方式",
    "标签",
    "创建时间",
    "解决时间",
    "关闭时间",
]

HEADER_ALIASES = {
    "工单号": ["工单号", "工单编号", "编号", "ticket_no", "ticketNo"],
    "标题": ["标题", "问题标题", "工单标题", "title"],
    "描述": ["描述", "问题描述", "现象", "工单描述", "description"],
    "状态": ["状态", "工单状态", "status"],
    "所属项目": [
        "所属项目",
        "项目",
        "所属商家",
        "商家",
        "客户",
        "project_name",
        "projectName",
        "merchant_name",
        "merchantName",
    ],
    "所属模块": ["所属模块", "模块", "系统模块", "module_name", "moduleName"],
    "问题分类": ["问题分类", "分类", "问题类型", "category_name", "categoryName"],
    "对方优先级": ["对方优先级", "客户优先级", "customer_priority", "customerPriority"],
    "内部优先级": ["内部优先级", "优先级", "internal_priority", "internalPriority"],
    "严重等级": ["严重等级", "严重程度", "severity"],
    "来源": ["来源", "工单来源", "source"],
    "提单人": ["提单人", "反馈人", "创建人", "reporter_name", "reporterName"],
    "当前处理人": ["当前处理人", "处理人", "负责人", "current_assignee_name", "currentAssigneeName"],
    "是否真实问题": ["是否真实问题", "是否问题", "真实问题", "is_problem", "isProblem"],
    "根因分类": ["根因分类", "原因分类", "root_cause_category", "rootCauseCategory"],
    "根因详情": ["根因详情", "原因", "根因", "root_cause", "rootCause", "root_cause_detail"],
    "触发原因": ["触发原因", "trigger_reason", "triggerReason"],
    "影响范围": ["影响范围", "impact_scope", "impactScope"],
    "复现步骤": ["复现步骤", "reproduce_steps", "reproduceSteps"],
    "排查过程": ["排查过程", "investigation_process", "investigationProcess"],
    "解决方案": ["解决方案", "修复方案", "solution", "fix_solution", "fixSolution"],
    "验证方式": ["验证方式", "verify_method", "verifyMethod"],
    "标签": ["标签", "tags"],
    "创建时间": ["创建时间", "提单时间", "create_time", "createTime"],
    "解决时间": ["解决时间", "resolved_at", "resolvedAt"],
    "关闭时间": ["关闭时间", "closed_at", "closedAt"],
}

STATUS_TEXT_MAP = {
    "待受理": TicketStatus.PENDING.value,
    "处理中": TicketStatus.PROCESSING.value,
    "待用户反馈": TicketStatus.WAIT_USER.value,
    "待开发": TicketStatus.WAIT_DEV.value,
    "待上线": TicketStatus.WAIT_RELEASE.value,
    "待验证": TicketStatus.WAIT_VERIFY.value,
    "已解决": TicketStatus.RESOLVED.value,
    "已关闭": TicketStatus.CLOSED.value,
    "已驳回": TicketStatus.REJECTED.value,
    "非问题": TicketStatus.NON_PROBLEM.value,
    "设计如此": TicketStatus.DESIGN_AS_EXPECTED.value,
    "用户误操作": TicketStatus.USER_MISOPERATION.value,
    "重复工单": TicketStatus.DUPLICATED.value,
}


def _current_user_name(current_user: CurrentUserModel) -> str:
    """
    获取当前登录用户名。
    :param current_user: 当前登录用户
    :return: 用户名
    """
    if not current_user or not current_user.user:
        return ""
    return current_user.user.user_name or current_user.user.nick_name or ""


def _current_user_id(current_user: CurrentUserModel) -> int | None:
    """
    获取当前登录用户ID。
    :param current_user: 当前登录用户
    :return: 用户ID
    """
    return current_user.user.user_id if current_user and current_user.user else None


def _ticket_no() -> str:
    """
    生成导入缺失工单号时使用的工单编号。
    :return: 工单编号
    """
    return f"TK{datetime.now().strftime('%Y%m%d')}{snowIdWorker.get_id()}"


class TicketImportService:
    """
    工单 Excel 导入服务。
    """

    @classmethod
    def build_import_template(cls) -> bytes:
        """
        生成工单导入模板。
        :return: Excel 文件字节
        """
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "工单导入模板"
        header_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
        for index, header in enumerate(IMPORT_HEADERS, 1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.fill = header_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            sheet.column_dimensions[cell.column_letter].width = 18

        samples = {
            "工单号": "FS-10001",
            "标题": "支付接口偶发超时",
            "描述": "客户反馈支付提交后页面长时间无响应",
            "状态": "已解决",
            "所属项目": "示例项目",
            "所属模块": "支付中心",
            "问题分类": "接口超时",
            "对方优先级": "P2",
            "内部优先级": "P1",
            "严重等级": "major",
            "来源": "客户反馈",
            "提单人": "张三",
            "当前处理人": "李四",
            "是否真实问题": "是",
            "根因分类": "代码缺陷",
            "根因详情": "下游重试未设置超时时间",
            "解决方案": "增加超时和熔断配置",
            "标签": "支付,超时",
            "创建时间": "2026-05-14 10:00:00",
            "解决时间": "2026-05-14 12:00:00",
        }
        sheet.append([samples.get(header, "") for header in IMPORT_HEADERS])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @classmethod
    async def import_excel(
        cls, query_db: Session, file_content: bytes, current_user: CurrentUserModel
    ) -> dict[str, Any]:
        """
        从 Excel 导入工单，重复工单号会跳过并返回提示。
        :param query_db: 数据库会话
        :param file_content: Excel 文件内容
        :param current_user: 当前登录用户
        :return: 导入汇总结果
        """
        workbook = load_workbook(filename=BytesIO(file_content), data_only=True)
        sheet = workbook.active
        header_map = cls._build_header_map([cell.value for cell in sheet[1]])
        rows = []
        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            row_data = cls._read_row(row, header_map)
            if not any(value not in (None, "") for value in row_data.values()):
                continue
            rows.append((row_index, row_data))

        incoming_nos = [cls._cell_text(row.get("工单号")) for _idx, row in rows if cls._cell_text(row.get("工单号"))]
        existing_nos = TicketDao.get_existing_ticket_nos(query_db, incoming_nos)
        seen_nos: set[str] = set()
        duplicate_nos = set(existing_nos)
        failed_rows = []
        imported_tickets: list[Ticket] = []
        now = datetime.now()

        try:
            for row_index, row in rows:
                ticket_no = cls._cell_text(row.get("工单号")) or _ticket_no()
                if ticket_no in existing_nos or ticket_no in seen_nos:
                    duplicate_nos.add(ticket_no)
                    continue
                seen_nos.add(ticket_no)
                title = cls._cell_text(row.get("标题"))
                if not title:
                    failed_rows.append({"row": row_index, "reason": "标题不能为空"})
                    continue
                relation_result = cls._resolve_import_relation_fields(query_db, row)
                if relation_result["message"]:
                    failed_rows.append({"row": row_index, "reason": relation_result["message"]})
                    continue

                create_time = cls._parse_datetime(row.get("创建时间")) or now
                resolved_at = cls._parse_datetime(row.get("解决时间"))
                closed_at = cls._parse_datetime(row.get("关闭时间"))
                status = cls._normalize_status(row.get("状态"))
                ticket = TicketDao.add_ticket(
                    query_db,
                    Ticket(
                        ticket_no=ticket_no,
                        title=title,
                        description=cls._cell_text(row.get("描述")),
                        project_id=relation_result["project_id"],
                        merchant_name=relation_result["project_name"],
                        module_id=relation_result["module_id"],
                        module_name=relation_result["module_name"],
                        category_name=cls._cell_text(row.get("问题分类")),
                        status=status,
                        customer_priority=cls._cell_text(row.get("对方优先级")) or "P3",
                        internal_priority=cls._cell_text(row.get("内部优先级")) or "P3",
                        severity=cls._cell_text(row.get("严重等级")),
                        source=cls._cell_text(row.get("来源")),
                        reporter_name=cls._cell_text(row.get("提单人")) or _current_user_name(current_user),
                        current_assignee_name=cls._cell_text(row.get("当前处理人")),
                        is_problem=cls._parse_bool(row.get("是否真实问题")),
                        root_cause=cls._cell_text(row.get("根因详情")),
                        solution=cls._cell_text(row.get("解决方案")),
                        resolved_at=resolved_at,
                        closed_at=closed_at,
                        total_process_seconds=cls._process_seconds(create_time, resolved_at, closed_at),
                        tags=cls._parse_tags(row.get("标签")),
                        extra_data={"import_source": "excel", "origin": "feishu_bitable"},
                        create_by=_current_user_name(current_user),
                        update_by=_current_user_name(current_user),
                        create_time=create_time,
                        update_time=now,
                    ),
                )
                TicketDao.add_status_history(
                    query_db,
                    TicketStatusHistory(
                        ticket_id=ticket.ticket_id,
                        from_status=None,
                        to_status=ticket.status,
                        operator_id=_current_user_id(current_user),
                        operator_name=_current_user_name(current_user),
                        started_at=create_time,
                        comment="Excel导入初始化状态",
                        create_time=create_time,
                    ),
                )
                TicketDao.add_event(
                    query_db,
                    TicketEvent(
                        ticket_id=ticket.ticket_id,
                        event_type=TicketEventType.TICKET_CREATED.value,
                        operator_id=_current_user_id(current_user),
                        operator_name=_current_user_name(current_user),
                        content="从飞书多维表格导出的 Excel 导入",
                        event_data={"ticket_no": ticket.ticket_no, "import_source": "excel"},
                        create_time=create_time,
                    ),
                )
                cls._save_rca(query_db, ticket, row, current_user, create_time)
                imported_tickets.append(ticket)

            embedding_count = TicketEmbeddingService.vectorize_tickets_for_scene(query_db, imported_tickets, "import")
            query_db.commit()
            return {
                "totalRows": len(rows),
                "importedCount": len(imported_tickets),
                "duplicateCount": len(duplicate_nos),
                "duplicateTicketNos": sorted(duplicate_nos),
                "failedCount": len(failed_rows),
                "failedRows": failed_rows,
                "embeddingCount": embedding_count,
            }
        except Exception:
            query_db.rollback()
            raise

    @classmethod
    def _build_header_map(cls, headers: list[Any]) -> dict[str, int]:
        """
        构建标准表头到列索引的映射。
        :param headers: Excel 第一行表头
        :return: 表头索引映射
        """
        normalized = {cls._cell_text(header).lower(): index for index, header in enumerate(headers)}
        header_map = {}
        for standard, aliases in HEADER_ALIASES.items():
            for alias in aliases:
                if alias.lower() in normalized:
                    header_map[standard] = normalized[alias.lower()]
                    break
        return header_map

    @classmethod
    def _resolve_import_relation_fields(cls, query_db: Session, row: dict[str, Any]) -> dict[str, Any]:
        """
        解析导入行中的测试项目和模块，并回填为工单关联字段。
        :param query_db: 数据库会话
        :param row: 导入行数据
        :return: 解析结果，包含项目模块ID、名称和失败信息
        """
        project_name = cls._cell_text(row.get("所属项目"))
        module_name = cls._cell_text(row.get("所属模块"))
        if not project_name:
            return {
                "message": "所属项目不能为空",
                "project_id": None,
                "project_name": "",
                "module_id": None,
                "module_name": "",
            }

        project = (
            query_db.query(HrmProject)
            .filter(
                HrmProject.project_name == project_name,
                HrmProject.status == QtrDataStatusEnum.normal.value,
                HrmProject.del_flag == "0",
            )
            .first()
        )
        if not project:
            return {
                "message": f"所属项目不存在或已停用：{project_name}",
                "project_id": None,
                "project_name": "",
                "module_id": None,
                "module_name": "",
            }

        module_id = None
        resolved_module_name = ""
        if module_name:
            module = (
                query_db.query(HrmModule)
                .filter(
                    HrmModule.module_name == module_name,
                    HrmModule.status == QtrDataStatusEnum.normal.value,
                    HrmModule.project_id == project.project_id,
                )
                .first()
            )
            if not module:
                return {
                    "message": f"所属模块不存在、已停用或不属于项目：{module_name}",
                    "project_id": project.project_id,
                    "project_name": project.project_name,
                    "module_id": None,
                    "module_name": "",
                }
            module_id = module.module_id
            resolved_module_name = module.module_name

        return {
            "message": "",
            "project_id": project.project_id,
            "project_name": project.project_name,
            "module_id": module_id,
            "module_name": resolved_module_name,
        }

    @classmethod
    def _read_row(cls, row: tuple[Any, ...], header_map: dict[str, int]) -> dict[str, Any]:
        """
        按标准字段读取 Excel 行。
        :param row: Excel 行数据
        :param header_map: 表头索引映射
        :return: 标准字段字典
        """
        return {
            header: row[index] if index < len(row) else None
            for header, index in header_map.items()
        }

    @classmethod
    def _save_rca(
        cls,
        query_db: Session,
        ticket: Ticket,
        row: dict[str, Any],
        current_user: CurrentUserModel,
        create_time: datetime,
    ) -> None:
        """
        保存导入行中的 RCA 字段。
        :param query_db: 数据库会话
        :param ticket: 工单对象
        :param row: 导入行数据
        :param current_user: 当前登录用户
        :param create_time: 创建时间
        :return: 无
        """
        has_rca = any(
            cls._cell_text(row.get(key))
            for key in ["根因分类", "根因详情", "触发原因", "影响范围", "复现步骤", "排查过程", "解决方案", "验证方式"]
        )
        if not has_rca:
            return
        TicketDao.upsert_rca(
            query_db,
            TicketRca(
                ticket_id=ticket.ticket_id,
                symptom=ticket.description,
                root_cause_category=cls._cell_text(row.get("根因分类")),
                root_cause_detail=cls._cell_text(row.get("根因详情")),
                trigger_reason=cls._cell_text(row.get("触发原因")),
                impact_scope=cls._cell_text(row.get("影响范围")),
                reproduce_steps=cls._cell_text(row.get("复现步骤")),
                investigation_process=cls._cell_text(row.get("排查过程")),
                fix_solution=cls._cell_text(row.get("解决方案")),
                verify_method=cls._cell_text(row.get("验证方式")),
                created_by_id=_current_user_id(current_user),
                created_by_name=_current_user_name(current_user),
                create_time=create_time,
            ),
        )

    @classmethod
    def _normalize_status(cls, value: Any) -> str:
        """
        将中文状态名或状态编码转换为系统状态编码。
        :param value: Excel 状态值
        :return: 状态编码
        """
        text = cls._cell_text(value)
        if not text:
            return TicketStatus.PENDING.value
        return STATUS_TEXT_MAP.get(text, text)

    @classmethod
    def _parse_tags(cls, value: Any) -> list[str]:
        """
        解析标签列。
        :param value: Excel 标签值
        :return: 标签列表
        """
        text = cls._cell_text(value)
        if not text:
            return []
        return [item.strip() for item in text.replace("，", ",").split(",") if item.strip()]

    @classmethod
    def _parse_bool(cls, value: Any) -> bool | None:
        """
        解析是否真实问题。
        :param value: Excel 单元格值
        :return: 布尔值或 None
        """
        text = cls._cell_text(value)
        if not text:
            return None
        if text in {"是", "真", "true", "True", "1", "真实问题"}:
            return True
        if text in {"否", "假", "false", "False", "0", "非问题"}:
            return False
        return None

    @classmethod
    def _parse_datetime(cls, value: Any) -> datetime | None:
        """
        解析 Excel 日期时间。
        :param value: Excel 单元格值
        :return: 日期时间或 None
        """
        if isinstance(value, datetime):
            return value
        text = cls._cell_text(value)
        if not text:
            return None
        for pattern in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(text[:19], pattern)
            except ValueError:
                continue
        return None

    @classmethod
    def _process_seconds(cls, create_time: datetime, resolved_at: datetime | None, closed_at: datetime | None) -> int:
        """
        根据导入时间计算处理耗时。
        :param create_time: 创建时间
        :param resolved_at: 解决时间
        :param closed_at: 关闭时间
        :return: 处理秒数
        """
        end_time = closed_at or resolved_at
        if not end_time:
            return 0
        return max(int((end_time - create_time).total_seconds()), 0)

    @classmethod
    def _cell_text(cls, value: Any) -> str:
        """
        将单元格值转换为去空白字符串。
        :param value: 单元格值
        :return: 字符串
        """
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return str(value).strip()
